import io
import json
import os
import re
from pathlib import Path

import openpyxl
import pytest
import qrcode
from fastapi.testclient import TestClient
from markupsafe import escape

from app import db, engine, hr, knowledge, llm, max_api, scenario, web
from app.config import cfg

PASSWORD = "test-pass-123"


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(cfg, "ADMIN_PASSWORD", PASSWORD)
    monkeypatch.setattr(cfg, "BOT_USERNAME", "korsik_test_bot")
    monkeypatch.setattr(cfg, "DEMO_MODE", False)
    monkeypatch.setattr(web, "LOGIN_DELAY", 0)
    # the model is off unless a test says otherwise
    monkeypatch.setattr(llm, "is_up", lambda: False)
    return TestClient(web.app, follow_redirects=False)


@pytest.fixture
def admin(client):
    assert client.post("/login", data={"password": PASSWORD}).status_code == 303
    return client


def add_recruiter(name="Ольга", max_user_id=None, code="olga"):
    return db.execute(
        "INSERT INTO recruiters (name, max_user_id, invite_code, created_at) VALUES (?, ?, ?, ?)",
        (name, max_user_id, code, db.now_iso()),
    )


def add_employee(name="Смирнова Анна", recruiter_id=None, tags="GrPP", max_user_id=None, code="anna",
                 tz="Europe/Moscow"):
    return db.execute(
        "INSERT INTO employees (full_name, first_name, start_date, timezone, tags, recruiter_id, invite_code,"
        " max_user_id, address, linked_at, created_at) VALUES (?, ?, '2026-11-02', ?, ?, ?, ?, ?, ?, ?, ?)",
        (name, name.split()[-1], tz, tags, recruiter_id, code, max_user_id, "ty" if max_user_id else None,
         db.now_iso() if max_user_id else None, db.now_iso()),
    )


SCENARIO_HEADERS = ["Шаг", "Когда", "После", "Теги", "Тип", "Текст на «ты»", "Текст на «вы»", "Кнопки", "Плохой ответ",
                    "Картинка", "Комментарий"]


def write_scenario(path, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Сценарий"
    sheet.append(SCENARIO_HEADERS)
    for row in rows:
        sheet.append(list(row) + [""] * (len(SCENARIO_HEADERS) - len(row)))
    workbook.save(path)
    return path


GOOD_ROWS = [
    ("10", "р1 10:00", "", "", "сообщение", "С первым днём, {имя}!"),
    ("20", "р3 17:30", "", "GTP, GrPP", "опрос", "Задачи понятны?", "", "Да | Нет", "Нет"),
]


@pytest.fixture
def scenario_file(tmp_path, monkeypatch):
    path = write_scenario(tmp_path / "scenario.xlsx", GOOD_ROWS)
    monkeypatch.setattr(cfg, "SCENARIO_PATH", str(path))
    return path


def add_step(employee_id, step_id, status="sent"):
    db.execute("INSERT INTO sent_steps (employee_id, step_id, status, sent_at) VALUES (?, ?, ?, ?)",
               (employee_id, step_id, status, db.now_iso()))


def add_message(employee_id, role, text):
    db.execute("INSERT INTO messages (employee_id, role, text, created_at) VALUES (?, ?, ?, ?)",
               (employee_id, role, text, db.now_iso()))


def add_answer(employee_id, step_id, question, answer, score=None):
    db.execute("INSERT INTO answers (employee_id, step_id, question, answer, score, created_at)"
               " VALUES (?, ?, ?, ?, ?, ?)", (employee_id, step_id, question, answer, score, db.now_iso()))


def add_pulse(employee_id, text, risk, reason=""):
    db.execute("INSERT INTO pulse (employee_id, text, tone, topic, risk, reason, created_at)"
               " VALUES (?, ?, 'негативный', 'другое', ?, ?, ?)", (employee_id, text, risk, reason, db.now_iso()))


# --- login


def test_pages_ask_for_login(client):
    for path in ("/", "/employees/1", "/recruiters", "/knowledge", "/questions", "/scenario", "/export.xlsx"):
        response = client.get(path)
        assert response.status_code == 303 and response.headers["location"] == "/login", path
    for path in ("/employees", "/recruiters", "/knowledge/upload", "/questions/1/answer", "/scenario/upload"):
        response = client.post(path, data={})
        assert response.status_code == 303 and response.headers["location"] == "/login", path
    assert client.get("/health").json() == {"ok": True}
    assert client.get("/login").status_code == 200


def test_wrong_password_does_not_let_in(client):
    for attempt in ("nope", "", PASSWORD + "x", "пароль"):
        assert client.post("/login", data={"password": attempt}).status_code == 401
    assert client.get("/").status_code == 303


def test_empty_admin_password_is_not_a_password(client, monkeypatch):
    # a missing ADMIN_PASSWORD in .env must lock the cabinet, not open it to an empty field
    monkeypatch.setattr(cfg, "ADMIN_PASSWORD", "")
    assert client.post("/login", data={"password": ""}).status_code == 401
    assert client.get("/").status_code == 303


def test_every_answer_carries_the_security_headers(client, admin):
    for response in (client.get("/health"), client.get("/"), client.get("/login"), client.post("/logout")):
        assert response.headers["cache-control"] == "no-store"
        assert response.headers["x-content-type-options"] == "nosniff"
        assert "default-src 'none'" in response.headers["content-security-policy"]


def test_login_gives_a_protected_cookie_and_logout_takes_it_back(client):
    response = client.post("/login", data={"password": PASSWORD})
    assert response.status_code == 303 and response.headers["location"] == "/"
    cookie = response.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=lax" in cookie
    assert PASSWORD not in response.headers["set-cookie"]
    assert client.get("/").status_code == 200
    assert client.post("/logout").status_code == 303
    assert client.get("/").status_code == 303


# --- employees and recruiters

NEW_EMPLOYEE = {"full_name": "Смирнова Анна Ивановна", "first_name": "", "start_date": "2026-11-02",
                "timezone": "Asia/Novosibirsk", "tags": "GrPP,  Luzhniki", "recruiter_id": ""}


def test_new_employee_gets_personal_link_and_qr(admin):
    recruiter_id = add_recruiter()
    response = admin.post("/employees", data={**NEW_EMPLOYEE, "recruiter_id": str(recruiter_id)})
    row = db.query_one("SELECT * FROM employees")
    assert response.status_code == 303 and response.headers["location"] == f"/employees/{row['id']}"
    assert (row["full_name"], row["first_name"], row["start_date"]) == ("Смирнова Анна Ивановна", "Анна", "2026-11-02")
    assert (row["timezone"], row["tags"], row["recruiter_id"]) == ("Asia/Novosibirsk", "GrPP, Luzhniki", recruiter_id)
    assert row["max_user_id"] is None

    link = f"https://max.ru/korsik_test_bot?start={row['invite_code']}"
    assert link in admin.get(response.headers["location"]).text
    qr = admin.get(f"/employees/{row['id']}/qr.png")
    expected = io.BytesIO()
    qrcode.make(link).save(expected, format="PNG")
    assert qr.headers["content-type"] == "image/png" and qr.content == expected.getvalue()


def test_every_employee_gets_a_different_code(admin):
    for _ in range(3):
        admin.post("/employees", data=NEW_EMPLOYEE)
    codes = [row["invite_code"] for row in db.query("SELECT invite_code FROM employees")]
    assert len(set(codes)) == 3 and all(len(code) >= 16 for code in codes)


def test_no_link_until_the_bot_name_is_known(admin, monkeypatch):
    monkeypatch.setattr(cfg, "BOT_USERNAME", "")
    admin.post("/employees", data=NEW_EMPLOYEE)
    employee_id = db.query_one("SELECT id FROM employees")["id"]
    assert "max.ru" not in admin.get(f"/employees/{employee_id}").text
    assert admin.get(f"/employees/{employee_id}/qr.png").status_code == 404


@pytest.mark.parametrize("field, value", [
    ("full_name", "   "),
    ("start_date", "2026-13-45"),
    ("start_date", ""),
    ("timezone", "Mars/Base"),
    ("recruiter_id", "999"),
    ("recruiter_id", "abc"),
])
def test_bad_employee_form_is_explained_not_saved(admin, field, value):
    response = admin.post("/employees", data={**NEW_EMPLOYEE, field: value})
    assert response.status_code == 303 and response.headers["location"] == "/"
    assert db.query("SELECT id FROM employees") == []
    assert "flash bad" in admin.get("/").text


def test_main_page_lists_employees_and_offers_russian_time_zones(admin):
    recruiter_id = add_recruiter("Ольга Петрова")
    add_employee("Иванов Пётр", recruiter_id=recruiter_id, tags="GTP", code="petr")
    page = admin.get("/").text
    assert "Иванов Пётр" in page and "Ольга Петрова" in page and "GTP" in page
    for zone in ("Europe/Kaliningrad", "Europe/Moscow", "Asia/Novosibirsk", "Asia/Kamchatka"):
        assert f'value="{zone}"' in page
    assert "Калининград" in page and "Камчатка" in page


def test_main_page_shows_the_model_mode(admin, monkeypatch):
    assert "упрощённый режим" in admin.get("/").text
    monkeypatch.setattr(llm, "is_up", lambda: True)
    page = admin.get("/").text
    assert "Модель работает" in page and "упрощённый режим" not in page


def test_unknown_employee_is_a_page_not_a_crash(admin):
    for path in ("/employees/999", "/employees/abc", "/no/such/page"):
        response = admin.get(path)
        assert response.status_code == 404 and "Такой страницы нет" in response.text, path


def test_recruiter_gets_a_link_that_the_bot_accepts(admin):
    assert admin.post("/recruiters", data={"name": "Ольга Петрова"}).status_code == 303
    row = db.query_one("SELECT * FROM recruiters")
    assert row["name"] == "Ольга Петрова" and row["max_user_id"] is None
    assert f"https://max.ru/korsik_test_bot?start=hr_{row['invite_code']}" in admin.get("/recruiters").text
    # the bot cuts «hr_» off the start payload and looks the rest up
    assert hr.link_recruiter(777, row["invite_code"])["name"] == "Ольга Петрова"
    assert "привязан" in admin.get("/recruiters").text
    assert admin.get(f"/recruiters/{row['id']}/qr.png").headers["content-type"] == "image/png"


def test_recruiter_without_a_name_is_refused(admin):
    admin.post("/recruiters", data={"name": "  "})
    assert db.query("SELECT id FROM recruiters") == []


def test_card_shows_steps_dialogue_answers_and_flags(admin, scenario_file):
    employee_id = add_employee(max_user_id=500)
    add_step(employee_id, "10", "sent")
    add_step(employee_id, "20", "answered")
    add_message(employee_id, "bot", "Добрый день, Анна!")
    add_message(employee_id, "employee", "Где столовая?")
    add_message(employee_id, "hr", "Столовая на втором этаже")
    add_answer(employee_id, "20", "Задачи понятны?", "Скорее да", 7)
    add_pulse(employee_id, "Хочу уйти", 2, "думает об уходе")
    add_pulse(employee_id, "Всё хорошо", 0)

    card = admin.get(f"/employees/{employee_id}").text
    assert "С первым днём, {имя}!" in card and "есть ответ" in card and "отправлен" in card
    assert "Добрый день, Анна!" in card and "Столовая на втором этаже" in card
    assert "Скорее да" in card and ">7<" in card
    assert "Хочу уйти" in card and "красный флаг" in card and "думает об уходе" in card
    assert "Всё хорошо" not in card


def test_demo_button_moves_time_and_sends_the_step(admin, monkeypatch):
    employee_id = add_employee(max_user_id=500)
    other_id = add_employee("Не открыл Ссылку", code="no-link")
    calls = []

    async def fake_demo_next(employee_id):
        calls.append(("demo_next", employee_id))
        return "12"

    async def fake_advance(employee_id):
        calls.append(("advance", employee_id))
        return 2

    monkeypatch.setattr(engine, "demo_next", fake_demo_next)
    monkeypatch.setattr(engine, "advance", fake_advance)

    # outside the demo mode there is no button and the address does nothing
    assert "Следующий шаг" not in admin.get(f"/employees/{employee_id}").text
    admin.post(f"/employees/{employee_id}/next")
    assert calls == []

    monkeypatch.setattr(cfg, "DEMO_MODE", True)
    assert "Следующий шаг" in admin.get(f"/employees/{employee_id}").text
    assert "Следующий шаг" not in admin.get(f"/employees/{other_id}").text
    response = admin.post(f"/employees/{employee_id}/next")
    assert response.headers["location"] == f"/employees/{employee_id}"
    assert calls == [("demo_next", employee_id), ("advance", employee_id)]
    assert "шагу 12" in admin.get(response.headers["location"]).text


def test_demo_step_that_cannot_be_sent_is_explained(admin, monkeypatch):
    employee_id = add_employee(max_user_id=500)
    monkeypatch.setattr(cfg, "DEMO_MODE", True)

    async def fake_demo_next(employee_id):
        return "12"

    async def broken_advance(employee_id):
        raise max_api.MaxError(503, "down")

    monkeypatch.setattr(engine, "demo_next", fake_demo_next)
    monkeypatch.setattr(engine, "advance", broken_advance)
    response = admin.post(f"/employees/{employee_id}/next")
    assert response.status_code == 303
    assert "сообщение не ушло" in admin.get(response.headers["location"]).text


# --- documents

VACATION = "Ежегодный оплачиваемый отпуск составляет 28 календарных дней.\n\nОтпуск оформляется заявлением."


def upload(admin, name="Отпуск.txt", content=VACATION.encode(), title="", tags=""):
    return admin.post("/knowledge/upload", data={"title": title, "tags": tags},
                      files={"file": (name, content, "application/octet-stream")})


def message_after(admin, response):
    assert response.status_code == 303
    return admin.get(response.headers["location"]).text


def uploads_dir():
    return Path(cfg.DATA_DIR) / "uploads"


def test_uploaded_document_is_listed_and_found_by_a_question(admin):
    page = message_after(admin, upload(admin, title="Правила отпуска", tags="GrPP"))
    assert "Правила отпуска" in page and "GrPP" in page and "flash bad" not in page

    # the stored file has a random name, never the one from the form
    saved = list(uploads_dir().iterdir())
    assert len(saved) == 1 and re.fullmatch(r"[0-9a-f]{32}\.txt", saved[0].name)
    assert db.query_one("SELECT filename FROM documents")["filename"] == saved[0].name

    found = admin.post("/knowledge/check", data={"question": "сколько дней отпуска", "tags": "GrPP"}).text
    assert "28 календарных дней" in found and "Правила отпуска" in found
    # a document with the tag of another site is not shown to this newcomer
    assert "28 календарных дней" not in admin.post("/knowledge/check",
                                                    data={"question": "сколько дней отпуска", "tags": "GTP"}).text
    assert "28 календарных дней" not in admin.post("/knowledge/check", data={"question": "сколько дней отпуска"}).text


def test_title_defaults_to_the_file_name(admin):
    upload(admin, name="Правила отпуска.md")
    assert db.query_one("SELECT title FROM documents")["title"] == "Правила отпуска"


def test_check_with_odd_characters_does_not_crash(admin):
    upload(admin)
    response = admin.post("/knowledge/check", data={"question": '")(* OR "', "tags": ""})
    assert response.status_code == 200 and "Ничего не нашлось" in response.text
    assert "Впишите вопрос" in message_after(admin, admin.post("/knowledge/check", data={"question": " "}))


@pytest.mark.parametrize("name", ["virus.exe", "script.html", "noextension", "photo.png"])
def test_unsupported_file_type_is_refused_with_a_reason(admin, name):
    page = message_after(admin, upload(admin, name=name))
    assert "flash bad" in page and "PDF, DOCX, PPTX, TXT или MD" in page
    assert db.query("SELECT id FROM documents") == [] and not uploads_dir().exists()


def test_file_over_20_mb_is_refused(admin):
    page = message_after(admin, upload(admin, content=b"a" * (20 * 1024 * 1024 + 1)))
    assert "flash bad" in page and "20 МБ" in page
    assert db.query("SELECT id FROM documents") == [] and not uploads_dir().exists()


def test_huge_request_is_stopped_before_it_is_read(admin):
    response = upload(admin, content=b"a" * (22 * 1024 * 1024))
    assert response.status_code == 413 and "20 МБ" in response.text
    assert not uploads_dir().exists()


def test_broken_and_empty_files_are_explained_not_a_server_error(admin):
    for name, content in (("Договор.docx", b"this is not a zip file"), ("Пусто.txt", b"   \n")):
        page = message_after(admin, upload(admin, name=name, content=content))
        assert "flash bad" in page, name
    assert db.query("SELECT id FROM documents") == []
    assert list(uploads_dir().iterdir()) == []


def test_deleted_document_is_gone_from_the_list_the_search_and_the_disk(admin):
    upload(admin, title="Правила отпуска", tags="GrPP")
    document_id = db.query_one("SELECT id FROM documents")["id"]
    page = message_after(admin, admin.post(f"/knowledge/{document_id}/delete"))
    assert "Правила отпуска" not in page
    assert db.query("SELECT id FROM chunks") == [] and list(uploads_dir().iterdir()) == []
    assert "28 календарных дней" not in admin.post("/knowledge/check",
                                                    data={"question": "сколько дней отпуска", "tags": "GrPP"}).text
    # a second click on an already deleted document is a message, not an error
    assert "flash bad" in message_after(admin, admin.post(f"/knowledge/{document_id}/delete"))


# --- questions

CONTEXT = [{"role": "employee", "text": "Привет"}, {"role": "bot", "text": "Здравствуйте!"}]


@pytest.fixture
def outbox(monkeypatch):
    sent = []

    async def fake_send(text, user_id=None, chat_id=None, buttons=None, image_token=None):
        sent.append({"user_id": user_id, "text": text})
        return {}

    monkeypatch.setattr(max_api, "send_message", fake_send)
    return sent


def add_question(employee_id, question="Где взять пропуск?", context=CONTEXT):
    return db.execute(
        "INSERT INTO escalations (employee_id, question, context, created_at) VALUES (?, ?, ?, ?)",
        (employee_id, question, json.dumps(context, ensure_ascii=False), db.now_iso()),
    )


def test_open_question_is_visible_even_without_a_recruiter(admin):
    employee_id = add_employee("Иванов Пётр", tags="GTP", max_user_id=500)
    add_question(employee_id)
    page = admin.get("/questions").text
    assert "Где взять пропуск?" in page and "Иванов Пётр" in page and "GTP" in page
    assert "Здравствуйте!" in page and "рекрутер не назначен" in page
    assert "Вопросы (1)" in page


def test_answer_from_the_cabinet_reaches_the_newcomer(admin, outbox):
    employee_id = add_employee(max_user_id=500)
    question_id = add_question(employee_id)
    response = admin.post(f"/questions/{question_id}/answer", data={"answer": "Пропуск в бюро пропусков, 2 этаж"})
    assert response.status_code == 303 and response.headers["location"] == "/questions"
    assert len(outbox) == 1 and outbox[0]["user_id"] == 500
    assert "Где взять пропуск?" in outbox[0]["text"] and "Пропуск в бюро пропусков, 2 этаж" in outbox[0]["text"]

    row = db.query_one("SELECT * FROM escalations WHERE id = ?", (question_id,))
    assert (row["status"], row["answer"], row["answered_by"]) == (
        "answered", "Пропуск в бюро пропусков, 2 этаж", "кабинет")
    assert db.query_one("SELECT role FROM messages WHERE employee_id = ?", (employee_id,))["role"] == "hr"
    page = admin.get("/questions").text
    assert "Вопросы (" not in page and "Открытых вопросов нет" in page
    assert "Пропуск в бюро пропусков, 2 этаж" in page


def test_second_answer_is_refused_with_who_answered(admin, outbox):
    question_id = add_question(add_employee(max_user_id=500))
    admin.post(f"/questions/{question_id}/answer", data={"answer": "Первый ответ"})
    response = admin.post(f"/questions/{question_id}/answer", data={"answer": "Второй ответ"})
    assert "На этот вопрос уже ответили (кабинет)" in admin.get(response.headers["location"]).text
    assert len(outbox) == 1


def test_missing_question_and_empty_answer_are_messages(admin, outbox):
    gone = admin.post("/questions/999/answer", data={"answer": "Привет"})
    assert hr.QUESTION_GONE in admin.get(gone.headers["location"]).text
    question_id = add_question(add_employee(max_user_id=500))
    empty = admin.post(f"/questions/{question_id}/answer", data={"answer": "   "})
    assert "Впишите ответ" in admin.get(empty.headers["location"]).text
    assert outbox == [] and db.query_one("SELECT status FROM escalations")["status"] == "open"


def test_answer_that_max_did_not_take_leaves_the_question_open(admin, monkeypatch):
    async def broken_send(*args, **kwargs):
        raise max_api.MaxError(503, "unavailable")

    monkeypatch.setattr(max_api, "send_message", broken_send)
    question_id = add_question(add_employee(max_user_id=500))
    response = admin.post(f"/questions/{question_id}/answer", data={"answer": "Ответ"})
    assert response.status_code == 303
    page = admin.get(response.headers["location"]).text
    assert "Не получилось отправить" in page and "Где взять пропуск?" in page
    assert db.query_one("SELECT status FROM escalations")["status"] == "open"


def test_question_of_a_newcomer_who_left_the_bot_cannot_be_answered(admin, outbox):
    question_id = add_question(add_employee(max_user_id=None))
    response = admin.post(f"/questions/{question_id}/answer", data={"answer": "Ответ"})
    assert "не привязан" in admin.get(response.headers["location"]).text
    assert outbox == [] and db.query_one("SELECT status FROM escalations")["status"] == "open"


# --- scenario

BAD_ROWS = [
    ("10", "когда-нибудь", "", "", "сообщение", "Текст"),
    ("11", "р1 10:00", "", "", "викторина", "Другой тип"),
    ("12", "р1 11:00", "", "", "сообщение", "я" * 4001),
]
NEW_ROWS = [("77", "старт", "", "", "сообщение", "Совсем новый сценарий")]


def upload_scenario(admin, rows, tmp_path):
    incoming = tmp_path / "incoming"
    incoming.mkdir(exist_ok=True)
    path = write_scenario(incoming / "new.xlsx", rows)
    return admin.post("/scenario/upload", files={"file": ("new.xlsx", path.read_bytes(), "application/octet-stream")})


@pytest.fixture
def reloads(monkeypatch):
    calls = []
    monkeypatch.setattr(engine, "reload_scenario", lambda: calls.append(1) or [])
    return calls


def test_scenario_with_errors_is_listed_and_does_not_replace_the_working_file(admin, scenario_file, reloads, tmp_path):
    before = scenario_file.read_bytes()
    expected = scenario.validate_scenario(write_scenario(tmp_path / "bad.xlsx", BAD_ROWS))
    assert len(expected) == 3

    response = upload_scenario(admin, BAD_ROWS, tmp_path)
    assert response.status_code == 400
    for problem in expected:
        assert str(escape(problem)) in response.text
    assert scenario_file.read_bytes() == before and reloads == []
    # no half-written copy is left next to the working file
    assert sorted(f.name for f in tmp_path.glob("*.xlsx")) == ["bad.xlsx", "scenario.xlsx"]
    assert list(tmp_path.glob(".new-*")) == []


def test_good_scenario_replaces_the_file_and_reloads(admin, scenario_file, reloads, tmp_path):
    before = scenario_file.read_bytes()
    response = upload_scenario(admin, NEW_ROWS, tmp_path)
    assert response.status_code == 303
    assert [step.id for step in scenario.load_scenario(scenario_file)] == ["77"]
    assert reloads == [1]
    # one step back stays on disk in case HR uploads the wrong version
    assert (Path(cfg.DATA_DIR) / "scenario-previous.xlsx").read_bytes() == before
    assert "flash bad" not in admin.get(response.headers["location"]).text


def test_scenario_upload_refuses_other_files(admin, scenario_file, reloads):
    before = scenario_file.read_bytes()
    notes = admin.post("/scenario/upload", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert notes.status_code == 303 and ".xlsx" in admin.get(notes.headers["location"]).text
    broken = admin.post("/scenario/upload", files={"file": ("scenario.xlsx", b"not a workbook", "text/plain")})
    assert broken.status_code == 400 and "не удалось открыть файл" in broken.text
    nothing = admin.post("/scenario/upload")
    assert nothing.status_code == 303
    assert scenario_file.read_bytes() == before and reloads == []


def test_scenario_page_reports_the_current_file_and_gives_it_back(admin, scenario_file):
    page = admin.get("/scenario").text
    assert "Шагов: 2" in page and "flash bad" not in page
    download = admin.get("/scenario/download")
    assert download.content == scenario_file.read_bytes()
    assert "attachment" in download.headers["content-disposition"]
    scenario_file.write_bytes(b"damaged")
    assert "не удалось открыть файл" in admin.get("/scenario").text


# --- excel

def seed_report():
    recruiter_id = add_recruiter("Ольга Петрова")
    employee_id = add_employee("Смирнова Анна", recruiter_id=recruiter_id, tags="GrPP", max_user_id=500)
    add_employee("Не Открыл", code="second", tags="GTP")
    add_step(employee_id, "10")
    add_answer(employee_id, "20", "Как дела?", "=1+1", 6)
    add_answer(employee_id, "21", "Что улучшить?", "Всё\x07 хорошо")
    add_pulse(employee_id, "Хочу уйти", 2, "причина")
    add_pulse(employee_id, "Трудно", 1, "причина")
    add_pulse(employee_id, "Всё хорошо", 0)
    add_question(employee_id, "Где взять пропуск?")
    answered_id = add_question(employee_id, "Когда зарплата?")
    db.execute("UPDATE escalations SET status = 'answered', answer = '5 и 20 числа', answered_by = 'Ольга Петрова'"
               " WHERE id = ?", (answered_id,))


def sheet_rows(workbook, title):
    rows = list(workbook[title].iter_rows(values_only=True))
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def test_excel_has_the_four_sheets_with_the_data(admin):
    seed_report()
    response = admin.get("/export.xlsx")
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in response.headers["content-disposition"]
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Новички", "Опросы", "Флаги", "Вопросы"]

    newcomers = {row["ФИО"]: row for row in sheet_rows(workbook, "Новички")}
    assert set(newcomers) == {"Смирнова Анна", "Не Открыл"}
    anna = newcomers["Смирнова Анна"]
    assert (anna["Рекрутер"], anna["Теги"], anna["Шагов отправлено"]) == ("Ольга Петрова", "GrPP", 1)
    assert (anna["Красных флагов"], anna["Открытых вопросов"]) == (1, 1)

    surveys = sheet_rows(workbook, "Опросы")
    assert [(row["Шаг"], row["Оценка"]) for row in surveys] == [("20", 6), ("21", None)]
    assert {row["Сообщение"] for row in sheet_rows(workbook, "Флаги")} == {"Хочу уйти", "Трудно"}
    questions = {row["Вопрос"]: row for row in sheet_rows(workbook, "Вопросы")}
    assert questions["Где взять пропуск?"]["Статус"] == "открыт"
    paid = questions["Когда зарплата?"]
    assert (paid["Статус"], paid["Ответ"]) == ("отвечен", "5 и 20 числа")


def test_excel_keeps_a_newcomers_text_from_running_as_a_formula(admin):
    seed_report()
    workbook = openpyxl.load_workbook(io.BytesIO(admin.get("/export.xlsx").content))
    cell = next(c for row in workbook["Опросы"].iter_rows(min_row=2) for c in row if c.value == "=1+1")
    assert cell.data_type == "s"
    # a control character in a chat message must not break the whole report
    assert {row["Ответ"] for row in sheet_rows(workbook, "Опросы")} == {"=1+1", "Всё хорошо"}


def test_empty_cabinet_still_gives_an_excel(admin):
    workbook = openpyxl.load_workbook(io.BytesIO(admin.get("/export.xlsx").content))
    assert workbook.sheetnames == ["Новички", "Опросы", "Флаги", "Вопросы"]


# --- escaping

SCRIPT = '<script>alert(1)</script>'


def test_user_text_is_escaped_on_every_page(admin, scenario_file, tmp_path):
    recruiter_id = add_recruiter(SCRIPT)
    employee_id = add_employee(SCRIPT, recruiter_id=recruiter_id, tags=SCRIPT, max_user_id=500)
    add_message(employee_id, "employee", SCRIPT)
    add_answer(employee_id, "20", SCRIPT, SCRIPT)
    add_pulse(employee_id, SCRIPT, 2, SCRIPT)
    add_question(employee_id, SCRIPT, [{"role": "employee", "text": SCRIPT}])
    answered_id = add_question(employee_id, "Старый вопрос")
    db.execute("UPDATE escalations SET status = 'answered', answer = ?, answered_by = ? WHERE id = ?",
               (SCRIPT, SCRIPT, answered_id))
    upload(admin, name=SCRIPT + ".txt", title=SCRIPT, tags=SCRIPT)
    write_scenario(scenario_file, [("10", "р1 10:00", "", "", "сообщение", SCRIPT)])
    add_step(employee_id, "10")

    pages = [admin.get(path) for path in ("/", f"/employees/{employee_id}", "/recruiters", "/knowledge", "/questions")]
    pages.append(admin.post("/knowledge/check", data={"question": '"' + SCRIPT, "tags": SCRIPT}))
    for response in pages:
        assert response.status_code == 200
        assert "<script" not in response.text and "&lt;script&gt;" in response.text, response.url

    # the list of scenario problems repeats what the file says
    bad = upload_scenario(admin, [("1", "р1 10:00", "", "", SCRIPT, "x")], tmp_path)
    assert bad.status_code == 400 and "<script" not in bad.text


def test_hr_answers_document_without_a_file_can_be_deleted(admin):
    knowledge.add_hr_answer("Где пропуск?", "На втором этаже", [])
    document_id = db.query_one("SELECT id FROM documents WHERE kind = 'hr_answers'")["id"]
    assert "Ответы HR" in admin.get("/knowledge").text
    page = message_after(admin, admin.post(f"/knowledge/{document_id}/delete"))
    assert "Ответы HR" not in page and "flash bad" not in page


def test_disk_errors_are_messages_not_a_server_error(admin, scenario_file, reloads, tmp_path, monkeypatch):
    # a data folder that cannot be created: its path is an ordinary file
    blocker = tmp_path / "blocker"
    blocker.write_text("x")
    monkeypatch.setattr(cfg, "DATA_DIR", str(blocker))
    assert "Не удалось сохранить файл" in message_after(admin, upload(admin))
    assert db.query("SELECT id FROM documents") == []

    def refuse(src, dst):
        raise PermissionError("read-only folder")

    before = scenario_file.read_bytes()
    monkeypatch.setattr(os, "replace", refuse)
    response = upload_scenario(admin, NEW_ROWS, tmp_path)
    assert "Не удалось записать файл сценария" in admin.get(response.headers["location"]).text
    assert scenario_file.read_bytes() == before and reloads == []
    assert list(tmp_path.glob(".new-*")) == []
