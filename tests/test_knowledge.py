import asyncio
from datetime import datetime, timedelta, timezone

import pymupdf
import pytest
import pytesseract
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

from app import db, knowledge, llm


@pytest.fixture(autouse=True)
def ollama_off(monkeypatch):
    """By default the model is switched off: every embed call fails like a sleeping Mac."""

    async def down(texts, **kwargs):
        raise llm.LLMUnavailable("off")

    monkeypatch.setattr(llm, "embed", down)
    monkeypatch.setattr(llm, "is_up", lambda: False)


@pytest.fixture
def ollama_on(monkeypatch):
    """A fake embedding: texts about food point one way, everything else the other way."""

    async def fake_embed(texts, **kwargs):
        food = ("столов", "поесть", "обед")
        return [[1.0, 0.0] if any(w in t.lower() for w in food) else [0.0, 1.0] for t in texts]

    monkeypatch.setattr(llm, "embed", fake_embed)
    monkeypatch.setattr(llm, "is_up", lambda: True)


def add(tmp_path, name, text, tags=(), title=None):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return asyncio.run(knowledge.add_document(path, title or name, list(tags)))


def find(query, tags=(), use_vectors=False):
    return asyncio.run(knowledge.search(query, list(tags), use_vectors=use_vectors))


PVTR_PART = """6.6. Рабочее время

6.6.1. Работникам Организации устанавливается пятидневная рабочая неделя с двумя
выходными днями. С понедельника по четверг: начало работы 9:00, окончание работы 18:00.

6.6.2. По согласованию с руководителем допускается иной режим рабочего времени.

- для работников с детьми до трех лет;

- для дистанционных работников.
6.6.3 В предпраздничные дни продолжительность рабочего времени уменьшается на один час.

6.6.4.Режим дистанционного работника устанавливается им по своему усмотрению."""


def test_split_chunks_cuts_by_clause_numbers():
    chunks = knowledge.split_chunks([("", PVTR_PART)])

    # OCR sometimes loses the space after the number: "6.6.4.Режим"
    assert [section for section, _ in chunks] == ["п. 6.6.1", "п. 6.6.2", "п. 6.6.3", "п. 6.6.4"]
    first, second, third, _ = (text for _, text in chunks)
    assert first.startswith("6.6. Рабочее время")
    assert "начало работы 9:00" in first
    assert "дистанционных работников" in second
    assert third.startswith("6.6.3 В предпраздничные")


def test_split_chunks_without_numbers_packs_paragraphs():
    paragraphs = [f"Абзац {i}. " + "слово " * 60 for i in range(6)]

    chunks = knowledge.split_chunks([("слайд 2", "\n\n".join(paragraphs))])

    assert len(chunks) > 1
    assert all(section == "слайд 2" for section, _ in chunks)
    assert all(len(text) <= 1100 for _, text in chunks)
    joined = " ".join(text for _, text in chunks)
    assert all(f"Абзац {i}." in joined for i in range(6))


def test_split_chunks_cuts_a_very_long_clause_and_keeps_its_number():
    long_clause = "6.3.3. Трудовой договор может быть расторгнут работодателем. " + "Основание описано подробно. " * 120

    chunks = knowledge.split_chunks([("", long_clause)])

    assert len(chunks) > 1
    assert {section for section, _ in chunks} == {"п. 6.3.3"}
    assert all(len(text) <= 1100 for _, text in chunks)


SMARTART = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>
<dgm:pt modelId="0" type="doc"><dgm:t><a:p/></dgm:t></dgm:pt>
<dgm:pt modelId="2"><dgm:t><a:p><a:r><a:t>Экология:</a:t></a:r></a:p></dgm:t></dgm:pt>
<dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>Качество </a:t></a:r><a:r><a:t>продукции</a:t></a:r></a:p></dgm:t></dgm:pt>
<dgm:pt modelId="3"><dgm:t><a:p><a:r><a:t>Беречь ресурсы</a:t></a:r></a:p></dgm:t></dgm:pt>
</dgm:ptLst><dgm:cxnLst>
<dgm:cxn modelId="c1" srcId="0" destId="1" srcOrd="0" destOrd="0"/>
<dgm:cxn modelId="c2" srcId="0" destId="2" srcOrd="1" destOrd="0"/>
<dgm:cxn modelId="c3" srcId="2" destId="3" srcOrd="0" destOrd="0"/>
</dgm:cxnLst></dgm:dataModel>"""


def test_extract_sections_reads_smartart_text(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Цели в области качества"
    diagram = Part(
        PackURI("/ppt/diagrams/data1.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
        prs.part.package,
        SMARTART.encode(),
    )
    slide.part.relate_to(diagram, RT.DIAGRAM_DATA)
    path = tmp_path / "policy.pptx"
    prs.save(path)

    sections = knowledge.extract_sections(path)

    assert [label for label, _ in sections] == ["слайд 1"]
    text = sections[0][1]
    assert "Цели в области качества" in text
    assert "Качество продукции" in text
    assert text.index("Качество продукции") < text.index("Экология:") < text.index("Беречь ресурсы")


def test_scan_without_tesseract_asks_for_a_text_version(tmp_path, monkeypatch):
    pdf = pymupdf.open()
    pdf.new_page()
    path = tmp_path / "scan.pdf"
    pdf.save(path)
    monkeypatch.setattr(pytesseract.pytesseract, "tesseract_cmd", str(tmp_path / "no-tesseract"))

    with pytest.raises(ValueError, match="текстов"):
        knowledge.extract_sections(path)


def test_search_keeps_to_employee_tags(tmp_path):
    add(tmp_path, "grpp.md", "Столовая на площадке GrPP работает с утра.", ["GrPP"])
    add(tmp_path, "gtp.md", "Столовая на площадке GTP закрыта на ремонт.", ["gtp"])
    add(tmp_path, "all.md", "Столовая для всех: правила для каждого.")

    titles = {hit.title for hit in find("столовая", ["GRPP"])}

    assert titles == {"grpp.md", "all.md"}
    assert {hit.title for hit in find("столовая", ["GTP"])} == {"gtp.md", "all.md"}
    assert {hit.title for hit in find("столовая", [])} == {"all.md"}


def test_documents_for_uses_the_same_tag_rule(tmp_path):
    add(tmp_path, "grpp.md", "Текст про GrPP.", ["GrPP", "офис"])
    add(tmp_path, "gtp.md", "Текст про GTP.", ["GTP"])
    add(tmp_path, "all.md", "Текст для всех.")

    assert {row["title"] for row in knowledge.documents_for(["Офис"])} == {"grpp.md", "all.md"}


KT_TEXT = """Записку можно опустить в ящик для обращений у проходной.

Сведения, составляющие коммерческую тайну, нельзя пересылать без шифрования."""


@pytest.mark.parametrize("use_vectors, stale_up", [(False, False), (True, False), (True, True)])
def test_keyword_search_finds_other_word_forms(tmp_path, monkeypatch, use_vectors, stale_up):
    add(tmp_path, "kt.txt", KT_TEXT)
    add(tmp_path, "other.txt", "Совсем другой текст про отпуск.")
    # stale_up: the watch loop has not noticed yet that Ollama went away
    monkeypatch.setattr(llm, "is_up", lambda: stale_up)

    for query, expected in [
        ("где ящики для обращения?", "ящик для обращений"),
        ("что такое коммерческая тайна", "коммерческую тайну"),
    ]:
        hits = find(query, use_vectors=use_vectors)
        assert hits, query
        assert expected in hits[0].text
        assert hits[0].keyword_rank == 1
        assert hits[0].vector_score is None


def test_vector_search_finds_meaning_without_shared_words(tmp_path, ollama_on):
    add(tmp_path, "food.md", "Столовая открыта в будни.")
    add(tmp_path, "pass.md", "Пропуск выдают на проходной.")

    hits = find("где можно поесть", use_vectors=True)

    assert hits[0].title == "food.md"
    assert hits[0].keyword_rank is None
    assert hits[0].vector_score == pytest.approx(1.0)


def test_embeddings_are_filled_later_when_the_model_returns(tmp_path, monkeypatch):
    add(tmp_path, "food.md", "Столовая открыта в будни.")
    assert db.query_one("SELECT COUNT(*) AS n FROM chunks WHERE embedding IS NULL")["n"] == 1

    async def fake_embed(texts, **kwargs):
        return [[0.5, 0.5] for _ in texts]

    monkeypatch.setattr(llm, "embed", fake_embed)

    assert asyncio.run(knowledge.embed_missing()) == 1
    assert db.query_one("SELECT COUNT(*) AS n FROM chunks WHERE embedding IS NULL")["n"] == 0


def test_failed_embedding_is_retried_after_a_pause_without_switching_the_mode(tmp_path, monkeypatch):
    add(tmp_path, "food.md", "Столовая открыта в будни.")
    tries = []

    async def timing_out(texts, mark_down=True):
        tries.append(mark_down)
        raise llm.LLMUnavailable("timeout")

    monkeypatch.setattr(llm, "embed", timing_out)
    assert asyncio.run(knowledge.embed_missing()) == 0
    assert asyncio.run(knowledge.embed_missing()) == 0
    assert tries == [False]

    later = (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat(timespec="seconds")
    monkeypatch.setattr(db, "now_iso", lambda: later)
    asyncio.run(knowledge.embed_missing())
    assert tries == [False, False]


def test_loading_the_same_file_again_replaces_the_document(tmp_path):
    add(tmp_path, "food.md", "Старое меню: борщ.")
    add(tmp_path, "food.md", "Новое меню: солянка.")

    assert len(db.query("SELECT id FROM documents")) == 1
    assert find("борщ") == []
    assert find("солянка")[0].text == "Новое меню: солянка."


def test_deleted_document_is_not_found(tmp_path):
    doc_id = add(tmp_path, "food.md", "Столовая открыта в будни.")

    knowledge.delete_document(doc_id)

    assert find("столовая") == []
    assert db.query("SELECT * FROM chunks") == []
    assert db.query("SELECT * FROM chunks_fts") == []


def test_hr_answer_goes_to_hr_answers_document_and_is_found():
    knowledge.add_hr_answer("Где взять пропуск?", "Пропуск выдают в первый день на проходной.", ["GTP"])
    knowledge.add_hr_answer("Есть ли парковка?", "Парковка есть у главного входа.", ["gtp"])

    hits = find("пропуска", ["GTP"])

    assert hits[0].title == "Ответы HR"
    assert "Пропуск выдают в первый день" in hits[0].text
    assert find("пропуска", ["GrPP"]) == []
    docs = [row for row in knowledge.documents_for(["GTP"]) if row["kind"] == "hr_answers"]
    assert len(docs) == 1
