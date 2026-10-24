import asyncio
import hmac
import io
import json
import logging
import os
import secrets
import shutil
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import jinja2
import openpyxl
import qrcode
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, RedirectResponse, Response
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from starlette.exceptions import HTTPException
from starlette.middleware.sessions import SessionMiddleware
from starlette.templating import Jinja2Templates

from app import db, engine, hours, hr, knowledge, llm, miniapp
from app.config import cfg
from app.scenario import ScenarioError, Step, load_scenario, validate_scenario

log = logging.getLogger(__name__)

LOGIN_DELAY = 1.0  # seconds of waiting after a wrong password, a cheap brake on guessing
SESSION_HOURS = 12
MAX_UPLOAD = 20 * 1024 * 1024
FORM_OVERHEAD = 1024 * 1024  # room for the other fields of an upload form
MAX_NAME = 200
UPLOAD_TYPES = (".pdf", ".docx", ".pptx", ".txt", ".md")
RISK_NAMES = {1: "есть трудность", 2: "красный флаг"}
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PUBLIC_PATHS = ("/login", "/health")
# pages have no scripts and load nothing from other hosts, so the policy can be strict
SECURITY_HEADERS = {
    "Cache-Control": "no-store",
    "X-Content-Type-Options": "nosniff",
    "Content-Security-Policy": "default-src 'none'; img-src 'self'; style-src 'unsafe-inline'; "
                               "form-action 'self'; base-uri 'none'; frame-ancestors 'none'",
}
# the MAX mini app runs MAX's bridge script and its own one, and lives inside a MAX frame
MINIAPP_HEADERS = SECURITY_HEADERS | {"Content-Security-Policy": miniapp.POLICY}
TIMEZONES = hours.TIMEZONES

# the admin is plain HTML pages, an auto-generated API page would only show the routes to strangers
app = FastAPI(docs_url=None, redoc_url=None, openapi_url=None)
app.include_router(miniapp.router)

templates = Jinja2Templates(env=jinja2.Environment(
    loader=jinja2.FileSystemLoader(Path(__file__).parent / "templates"),
    autoescape=True,
))


def local_time(stamp: str | None) -> str:
    """UTC stamps from the database, shown in the cabinet's own time zone."""
    if not stamp:
        return ""
    try:
        return datetime.fromisoformat(stamp).astimezone(ZoneInfo(cfg.DEFAULT_TZ)).strftime("%d.%m.%Y %H:%M")
    except (ValueError, KeyError):  # a bad stamp, or a DEFAULT_TZ that is not a time zone
        return stamp


templates.env.filters["local"] = local_time
templates.env.filters["zone"] = lambda name: TIMEZONES.get(name, name)
templates.env.globals["tz_name"] = cfg.DEFAULT_TZ
templates.env.globals["open_questions"] = lambda: db.query_one(
    "SELECT COUNT(*) AS n FROM escalations WHERE status = 'open'")["n"]
templates.env.globals["overdue_questions"] = hr.overdue_count
templates.env.globals["day_names"] = hours.DAY_NAMES
templates.env.filters["schedule"] = lambda row: hours.from_row(row).describe()


@app.middleware("http")
async def guard(request: Request, call_next):
    size = request.headers.get("content-length", "")
    # the mini app has no cabinet session: each of its requests carries MAX's signed initData instead
    in_miniapp = request.url.path == "/app" or request.url.path.startswith("/app/")
    if size.isdigit() and int(size) > MAX_UPLOAD + FORM_OVERHEAD:
        # answered before the body is read, so a huge upload never lands on disk
        response = page(request, "message.html", 413, text="Файл больше 20 МБ.")
    elif in_miniapp or request.url.path in PUBLIC_PATHS or request.session.get("admin"):
        response = await call_next(request)
    else:
        response = RedirectResponse("/login", status_code=303)
    response.headers.update(MINIAPP_HEADERS if in_miniapp else SECURITY_HEADERS)
    return response


if not cfg.SECRET_KEY:
    log.warning("SECRET_KEY is empty: sessions use a random key and end with every restart")
if not cfg.ADMIN_PASSWORD:
    log.warning("ADMIN_PASSWORD is empty: nobody can log in to the cabinet")
# Added after guard, so it wraps guard and the session is ready when guard looks at it.
# No Secure flag: the cabinet is also opened over plain http on 127.0.0.1, the HTTPS proxy sits in front of it.
app.add_middleware(
    SessionMiddleware,
    secret_key=cfg.SECRET_KEY or secrets.token_hex(32),
    session_cookie="korsik_session",
    max_age=SESSION_HOURS * 3600,
    same_site="lax",
)


def page(request: Request, name: str, status: int = 200, **context):
    context["flash"] = request.session.pop("flash", None)
    return templates.TemplateResponse(request, name, context, status_code=status)


def back(request: Request, url: str, message: str | None = None, bad: bool = False):
    """Redirect after a POST; the message shows once on the next page."""
    if message:
        request.session["flash"] = {"text": message[:300], "bad": bad}
    return RedirectResponse(url, status_code=303)


@app.exception_handler(HTTPException)
@app.exception_handler(RequestValidationError)
async def http_error(request: Request, error: Exception):
    # a mistyped address like /employees/abc is also "no such page" for a person
    status = getattr(error, "status_code", 404)
    text = "Такой страницы нет." if status == 404 else "Запрос не получился."
    return page(request, "message.html", status, text=text)


@app.get("/health")
def health():
    return {"ok": True}


@app.get("/login")
def login_form(request: Request):
    return page(request, "login.html")


@app.post("/login")
async def login(request: Request, password: str = Form("")):
    # bytes, because compare_digest refuses str with non-ASCII letters; an empty setting never matches
    if cfg.ADMIN_PASSWORD and hmac.compare_digest(password.encode(), cfg.ADMIN_PASSWORD.encode()):
        request.session.clear()
        request.session["admin"] = True
        return RedirectResponse("/", status_code=303)
    await asyncio.sleep(LOGIN_DELAY)
    return page(request, "login.html", 401, error="Неверный пароль.")


@app.post("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/")
def index(request: Request):
    employees = db.query(
        "SELECT e.*, r.name AS recruiter_name, w.name AS workplace_name,"
        " (SELECT COUNT(*) FROM sent_steps s WHERE s.employee_id = e.id AND s.status != 'skipped') AS steps_sent,"
        " (SELECT COUNT(*) FROM escalations q WHERE q.employee_id = e.id AND q.status = 'open') AS open_questions,"
        " (SELECT COUNT(*) FROM pulse p WHERE p.employee_id = e.id AND p.risk = 2) AS red_flags"
        " FROM employees e LEFT JOIN recruiters r ON r.id = e.recruiter_id"
        " LEFT JOIN workplaces w ON w.id = e.workplace_id ORDER BY e.start_date DESC, e.id DESC"
    )
    recruiters = db.query("SELECT id, name FROM recruiters ORDER BY name")
    default_zone = cfg.DEFAULT_TZ if cfg.DEFAULT_TZ in TIMEZONES else "Europe/Moscow"
    return page(request, "index.html", employees=employees, recruiters=recruiters, zones=TIMEZONES,
                default_zone=default_zone, tags=known_tags(), model_up=llm.is_up(), workplaces=list_workplaces())


def list_workplaces():
    return db.query(
        "SELECT w.*, (SELECT COUNT(*) FROM employees e WHERE e.workplace_id = w.id) AS employees,"
        " (SELECT group_concat(r.name, ', ') FROM recruiter_workplaces rw JOIN recruiters r ON r.id = rw.recruiter_id"
        "  WHERE rw.workplace_id = w.id) AS recruiter_names"
        " FROM workplaces w ORDER BY w.name")


def read_schedule(tz: str, days: list[str], start: str, end: str) -> tuple[tuple | None, str]:
    return hours.read_form(tz, days, start, end)


def known_tags() -> list[str]:
    """Tags already used by the scenario and the documents: a typo in a tag silently hides steps from a newcomer."""
    found = {}
    for step in read_scenario()[0]:
        for tag in step.tags:
            found.setdefault(tag.lstrip("-").strip().lower(), tag.lstrip("-").strip())
    for row in db.query("SELECT tags FROM documents"):
        for tag in db.split_tags(row["tags"]):
            found.setdefault(tag.lower(), tag)
    return sorted(found.values(), key=str.lower)


@app.post("/employees")
def add_employee(request: Request, full_name: str = Form(""), first_name: str = Form(""), start_date: str = Form(""),
                 tz: str = Form("", alias="timezone"), tags: str = Form(""), recruiter_id: str = Form(""),
                 workplace_id: str = Form("")):
    full_name = " ".join(full_name.split())
    workplace = None
    if workplace_id.strip():
        workplace = db.query_one("SELECT * FROM workplaces WHERE id = ?",
                                 (int(workplace_id),)) if workplace_id.strip().isdigit() else None
        if workplace is None:
            return back(request, "/", "Такого места работы нет.", bad=True)
    elif db.query_one("SELECT 1 FROM workplaces"):
        return back(request, "/", "Выберите место работы: от него зависят теги, часовой пояс и рабочие часы.",
                    bad=True)
    if workplace is not None:
        # the workplace brings its tags and zone; extra tags typed by HR are kept on top
        if tz in ("", "workplace"):
            tz = workplace["timezone"]
        extra = [t for t in db.split_tags(tags) if t.lower() not in {w.lower() for w in db.split_tags(workplace["tags"])}]
        tags = ", ".join(db.split_tags(workplace["tags"]) + extra)
        if not recruiter_id.strip():
            # the first recruiter who covers this workplace takes the newcomer
            owner = db.query_one("SELECT recruiter_id FROM recruiter_workplaces WHERE workplace_id = ?"
                                 " ORDER BY recruiter_id LIMIT 1", (workplace["id"],))
            recruiter_id = str(owner["recruiter_id"]) if owner else ""
    names = full_name.split()
    tags = ", ".join(db.split_tags(tags))
    try:
        start = date.fromisoformat(start_date)
    except ValueError:
        start = None
    recruiter_id = recruiter_id.strip()

    if not names or len(full_name) > MAX_NAME:
        return back(request, "/", "Впишите ФИО, не длиннее 200 знаков.", bad=True)
    if start is None or not 2020 <= start.year <= 2100:
        return back(request, "/", "Дата выхода нужна в виде год-месяц-день, например 2026-11-02.", bad=True)
    if tz not in TIMEZONES:
        return back(request, "/", "Выберите часовой пояс из списка.", bad=True)
    if len(tags) > MAX_NAME or len(first_name) > MAX_NAME:
        return back(request, "/", "Теги и имя слишком длинные.", bad=True)
    if recruiter_id and not (recruiter_id.isdigit()
                             and db.query_one("SELECT 1 FROM recruiters WHERE id = ?", (int(recruiter_id),))):
        return back(request, "/", "Такого рекрутера нет.", bad=True)

    # «Фамилия Имя Отчество»: the name to address is the second word unless HR typed it separately
    first_name = first_name.strip() or (names[1] if len(names) > 1 else names[0])
    employee_id = db.execute(
        "INSERT INTO employees (full_name, first_name, start_date, timezone, tags, recruiter_id, invite_code,"
        " created_at, workplace_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (full_name, first_name, start.isoformat(), tz, tags, int(recruiter_id) if recruiter_id else None,
         secrets.token_hex(8), db.now_iso(), workplace["id"] if workplace else None),
    )
    return back(request, f"/employees/{employee_id}", "Новичок добавлен. Отправьте ему ссылку или QR.")


@app.get("/employees/{employee_id}")
def employee_card(request: Request, employee_id: int):
    employee = db.query_one(
        "SELECT e.*, r.name AS recruiter_name, w.name AS workplace_name FROM employees e"
        " LEFT JOIN recruiters r ON r.id = e.recruiter_id LEFT JOIN workplaces w ON w.id = e.workplace_id"
        " WHERE e.id = ?", (employee_id,))
    if employee is None:
        raise HTTPException(404)
    schedule = hours.of_employee(employee)
    steps = db.query("SELECT * FROM sent_steps WHERE employee_id = ? ORDER BY sent_at, rowid", (employee_id,))
    titles = {step.id: step.text_ty for step in read_scenario()[0]}
    # the newest 300 are enough to follow a dialogue, three months of chat would make the page heavy
    messages = db.query("SELECT * FROM (SELECT * FROM messages WHERE employee_id = ? ORDER BY id DESC LIMIT 300)"
                        " ORDER BY id", (employee_id,))
    answers = db.query("SELECT * FROM answers WHERE employee_id = ? ORDER BY id", (employee_id,))
    flags = db.query("SELECT * FROM pulse WHERE employee_id = ? AND risk >= 1 ORDER BY id", (employee_id,))
    now = engine.employee_now(employee).astimezone(ZoneInfo(employee["timezone"]))
    return page(request, "employee.html", employee=employee, link=engine.invite_link(employee["invite_code"]),
                steps=steps, titles=titles, messages=messages, answers=answers, flags=flags, demo=cfg.DEMO_MODE,
                employee_now=now.strftime("%d.%m.%Y %H:%M"), workplaces=list_workplaces(),
                recruiters=db.query("SELECT id, name FROM recruiters ORDER BY name"),
                schedule=schedule.describe() if schedule else "", late=engine.is_late(employee))


@app.post("/employees/{employee_id}/move")
def move_employee(request: Request, employee_id: int, workplace_id: str = Form(""), recruiter_id: str = Form("")):
    url = f"/employees/{employee_id}"
    employee = db.query_one("SELECT * FROM employees WHERE id = ?", (employee_id,))
    if employee is None:
        raise HTTPException(404)
    if workplace_id and not (workplace_id.isdigit()
                             and db.query_one("SELECT 1 FROM workplaces WHERE id = ?", (int(workplace_id),))):
        return back(request, url, "Такого места работы нет.", bad=True)
    if recruiter_id and not (recruiter_id.isdigit()
                             and db.query_one("SELECT 1 FROM recruiters WHERE id = ?", (int(recruiter_id),))):
        return back(request, url, "Такого рекрутера нет.", bad=True)
    new_workplace = int(workplace_id) if workplace_id else None
    if new_workplace != employee["workplace_id"]:
        hr.change_workplace(employee_id, new_workplace)
    db.execute("UPDATE employees SET recruiter_id = ? WHERE id = ?",
               (int(recruiter_id) if recruiter_id else None, employee_id))
    return back(request, url, "Сохранено. Следующие шаги и ответы пойдут по новому месту работы.")


@app.post("/employees/{employee_id}/next")
async def next_step(request: Request, employee_id: int):
    url = f"/employees/{employee_id}"
    if not cfg.DEMO_MODE:
        return back(request, url, "Кнопка работает только в демо-режиме.", bad=True)
    step_id = await engine.demo_next(employee_id)
    if step_id is None:
        return back(request, url, "Дальше по сценарию шагов нет, или новичок ещё не открыл ссылку.", bad=True)
    try:
        sent = await engine.advance(employee_id)
    except Exception as e:
        log.warning("employee %s: demo step was not sent: %s", employee_id, type(e).__name__)
        return back(request, url, f"Время сдвинуто к шагу {step_id}, но сообщение не ушло. "
                                  "Планировщик отправит его сам.", bad=True)
    return back(request, url, f"Время сдвинуто к шагу {step_id}, отправлено сообщений: {sent}.")


@app.get("/employees/{employee_id}/qr.png")
def employee_qr(employee_id: int):
    employee = db.query_one("SELECT invite_code FROM employees WHERE id = ?", (employee_id,))
    return qr_response(engine.invite_link(employee["invite_code"]) if employee else None)


def qr_response(link: str | None) -> Response:
    if link is None:
        raise HTTPException(404)
    image = io.BytesIO()
    qrcode.make(link).save(image, format="PNG")
    return Response(image.getvalue(), media_type="image/png")


@app.get("/recruiters")
def recruiters_page(request: Request):
    rows = db.query(
        "SELECT r.*, (SELECT COUNT(*) FROM employees e WHERE e.recruiter_id = r.id) AS employees,"
        " (SELECT group_concat(w.name, ', ') FROM recruiter_workplaces rw JOIN workplaces w ON w.id = rw.workplace_id"
        "  WHERE rw.recruiter_id = r.id) AS workplace_names"
        " FROM recruiters r ORDER BY r.id"
    )
    links = {row["id"]: engine.invite_link("hr_" + row["invite_code"]) for row in rows}
    now = datetime.now(ZoneInfo("UTC"))
    working = {row["id"]: hours.of_recruiter(row).is_open(now, engine.holidays()) for row in rows}
    stats = {row["id"]: hr.recruiter_stats(row["id"]) for row in rows}
    return page(request, "recruiters.html", recruiters=rows, links=links, working=working, stats=stats,
                zones=TIMEZONES, workplaces=list_workplaces(), default_zone=cfg.DEFAULT_TZ)


@app.post("/recruiters")
def add_recruiter(request: Request, name: str = Form(""), tz: str = Form("Europe/Moscow", alias="timezone"),
                  days: list[str] = Form([]), work_start: str = Form("09:00"), work_end: str = Form("18:00"),
                  workplaces: list[str] = Form([])):
    name = " ".join(name.split())
    if not name or len(name) > MAX_NAME:
        return back(request, "/recruiters", "Впишите имя рекрутера.", bad=True)
    schedule, problem = read_schedule(tz, days, work_start, work_end)
    if schedule is None:
        return back(request, "/recruiters", problem, bad=True)
    recruiter_id = db.execute(
        "INSERT INTO recruiters (name, invite_code, created_at, timezone, work_days, work_start, work_end)"
        " VALUES (?, ?, ?, ?, ?, ?, ?)", (name, secrets.token_hex(8), db.now_iso(), *schedule))
    save_recruiter_workplaces(recruiter_id, workplaces)
    return back(request, "/recruiters", "Рекрутер добавлен. Пусть откроет ссылку в МАКС.")


def save_recruiter_workplaces(recruiter_id: int, workplaces: list[str]) -> None:
    known = {row["id"] for row in db.query("SELECT id FROM workplaces")}
    picked = {int(w) for w in workplaces if w.isdigit() and int(w) in known}
    conn = db.connect()
    try:
        with conn:
            conn.execute("DELETE FROM recruiter_workplaces WHERE recruiter_id = ?", (recruiter_id,))
            conn.executemany("INSERT INTO recruiter_workplaces (recruiter_id, workplace_id) VALUES (?, ?)",
                             [(recruiter_id, w) for w in sorted(picked)])
    finally:
        conn.close()


@app.get("/recruiters/{recruiter_id}")
def recruiter_card(request: Request, recruiter_id: int):
    recruiter = db.query_one("SELECT * FROM recruiters WHERE id = ?", (recruiter_id,))
    if recruiter is None:
        raise HTTPException(404)
    mine = {row["workplace_id"] for row in db.query(
        "SELECT workplace_id FROM recruiter_workplaces WHERE recruiter_id = ?", (recruiter_id,))}
    employees = db.query(
        "SELECT e.*, w.name AS workplace_name FROM employees e LEFT JOIN workplaces w ON w.id = e.workplace_id"
        " WHERE e.recruiter_id = ? ORDER BY e.start_date DESC", (recruiter_id,))
    questions = db.query(
        "SELECT q.*, e.full_name FROM escalations q JOIN employees e ON e.id = q.employee_id"
        " WHERE e.recruiter_id = ? ORDER BY q.id DESC LIMIT 100", (recruiter_id,))
    now = datetime.now(ZoneInfo("UTC"))
    states = {q["id"]: hr.question_state(q, now) for q in questions if q["status"] == "open"}
    took = {q["id"]: hours.waited_text(q["created_at"], datetime.fromisoformat(q["answered_at"]))
            for q in questions if q["answered_at"]}
    schedule = hours.of_recruiter(recruiter)
    return page(request, "recruiter.html", recruiter=recruiter, stats=hr.recruiter_stats(recruiter_id),
                employees=employees, questions=questions, states=states, took=took, zones=TIMEZONES,
                workplaces=list_workplaces(), mine=mine, days=hours.parse_days(recruiter["work_days"]),
                working=schedule.is_open(now, engine.holidays()), schedule=schedule.describe(),
                link=engine.invite_link("hr_" + recruiter["invite_code"]))


@app.post("/recruiters/{recruiter_id}")
def update_recruiter(request: Request, recruiter_id: int, name: str = Form(""),
                     tz: str = Form("Europe/Moscow", alias="timezone"), days: list[str] = Form([]),
                     work_start: str = Form(""), work_end: str = Form(""), workplaces: list[str] = Form([])):
    url = f"/recruiters/{recruiter_id}"
    if db.query_one("SELECT 1 FROM recruiters WHERE id = ?", (recruiter_id,)) is None:
        raise HTTPException(404)
    name = " ".join(name.split())
    if not name or len(name) > MAX_NAME:
        return back(request, url, "Впишите имя рекрутера.", bad=True)
    schedule, problem = read_schedule(tz, days, work_start, work_end)
    if schedule is None:
        return back(request, url, problem, bad=True)
    db.execute("UPDATE recruiters SET name = ?, timezone = ?, work_days = ?, work_start = ?, work_end = ? WHERE id = ?",
               (name, *schedule, recruiter_id))
    save_recruiter_workplaces(recruiter_id, workplaces)
    return back(request, url, "Сохранено.")


@app.get("/workplaces")
def workplaces_page(request: Request):
    return page(request, "workplaces.html", workplaces=list_workplaces(), zones=TIMEZONES, tags=known_tags(),
                default_zone=cfg.DEFAULT_TZ)


@app.post("/workplaces")
def save_workplace(request: Request, workplace_id: str = Form(""), name: str = Form(""), tags: str = Form(""),
                   tz: str = Form("", alias="timezone"), days: list[str] = Form([]), work_start: str = Form(""),
                   work_end: str = Form("")):
    name = " ".join(name.split())
    tags = ", ".join(db.split_tags(tags))
    if not name or len(name) > MAX_NAME or len(tags) > MAX_NAME:
        return back(request, "/workplaces", "Впишите название места работы, не длиннее 200 знаков.", bad=True)
    schedule, problem = read_schedule(tz, days, work_start, work_end)
    if schedule is None:
        return back(request, "/workplaces", problem, bad=True)
    same = db.query_one("SELECT id FROM workplaces WHERE lower(name) = lower(?)", (name,))
    if workplace_id.isdigit():
        if same and same["id"] != int(workplace_id):
            return back(request, "/workplaces", "Место работы с таким названием уже есть.", bad=True)
        old = db.query_one("SELECT * FROM workplaces WHERE id = ?", (int(workplace_id),))
        if old is None:
            raise HTTPException(404)
        db.execute("UPDATE workplaces SET name = ?, tags = ?, timezone = ?, work_days = ?, work_start = ?,"
                   " work_end = ? WHERE id = ?", (name, tags, *schedule, old["id"]))
        if old["tags"] != tags or old["timezone"] != schedule[0]:
            # newcomers of this workplace follow its new tags and zone; tags HR typed by hand for them stay
            old_tags = {t.lower() for t in db.split_tags(old["tags"])}
            for employee in db.query("SELECT * FROM employees WHERE workplace_id = ?", (old["id"],)):
                own = [t for t in db.split_tags(employee["tags"]) if t.lower() not in old_tags]
                merged = db.split_tags(tags) + [t for t in own if t.lower() not in {x.lower() for x in db.split_tags(tags)}]
                db.execute("UPDATE employees SET tags = ?, timezone = ? WHERE id = ?",
                           (", ".join(merged), schedule[0], employee["id"]))
        return back(request, "/workplaces", "Место работы сохранено.")
    if same:
        return back(request, "/workplaces", "Место работы с таким названием уже есть.", bad=True)
    db.execute("INSERT INTO workplaces (name, tags, timezone, work_days, work_start, work_end, created_at)"
               " VALUES (?, ?, ?, ?, ?, ?, ?)", (name, tags, *schedule, db.now_iso()))
    return back(request, "/workplaces", "Место работы добавлено.")


@app.post("/workplaces/{workplace_id}/delete")
def delete_workplace(request: Request, workplace_id: int):
    if db.query_one("SELECT 1 FROM employees WHERE workplace_id = ?", (workplace_id,)):
        return back(request, "/workplaces", "Здесь есть новички: сначала переведите их на другое место работы.",
                    bad=True)
    db.execute("DELETE FROM workplaces WHERE id = ?", (workplace_id,))
    return back(request, "/workplaces", "Место работы удалено.")


@app.get("/dashboard")
def dashboard(request: Request):
    now = datetime.now(ZoneInfo("UTC"))
    week_ago = (now - timedelta(days=7)).isoformat(timespec="seconds")
    employees = db.query("SELECT * FROM employees")
    stages = {"before": 0, "onboarding": 0, "done": 0, "not_linked": 0}
    for e in employees:
        if not e["max_user_id"]:
            stages["not_linked"] += 1
            continue
        today = engine.employee_now(e).astimezone(ZoneInfo(e["timezone"])).date()
        start = date.fromisoformat(e["start_date"])
        months = (today.year - start.year) * 12 + today.month - start.month - (today.day < start.day)
        stages["before" if today < start else "done" if months >= cfg.PROBATION_MONTHS else "onboarding"] += 1
    opened = db.query("SELECT q.*, e.full_name, e.workplace_id, e.recruiter_id FROM escalations q"
                      " JOIN employees e ON e.id = q.employee_id WHERE q.status = 'open' ORDER BY q.id")
    states = {q["id"]: hr.question_state(q, now) for q in opened}
    places = []
    for w in list(list_workplaces()) + [None]:
        wid = w["id"] if w else None
        mine = [e for e in employees if e["workplace_id"] == wid]
        if w is None and not mine:
            continue
        ids = {e["id"] for e in mine}
        places.append({
            "workplace": w, "employees": len(mine),
            "open": sum(1 for q in opened if q["employee_id"] in ids),
            "overdue": sum(1 for q in opened if q["employee_id"] in ids and states[q["id"]]["overdue"]),
            "flags": db.query_one(
                f"SELECT COUNT(*) AS n FROM pulse WHERE risk = 2 AND created_at >= ? AND employee_id IN "
                f"({','.join('?' * len(ids)) or 'NULL'})", (week_ago, *ids))["n"],
        })
    recruiters = []
    for r in db.query("SELECT * FROM recruiters ORDER BY name"):
        recruiters.append({"recruiter": r, "stats": hr.recruiter_stats(r["id"]),
                           "working": hours.of_recruiter(r).is_open(now, engine.holidays()),
                           "employees": sum(1 for e in employees if e["recruiter_id"] == r["id"])})
    red = db.query("SELECT p.*, e.full_name FROM pulse p JOIN employees e ON e.id = p.employee_id"
                   " WHERE p.risk = 2 AND p.created_at >= ? ORDER BY p.id DESC LIMIT 10", (week_ago,))
    topics = db.query("SELECT topic, COUNT(*) AS n FROM pulse WHERE risk >= 1 AND topic != 'другое'"
                      " GROUP BY topic ORDER BY n DESC LIMIT 5")
    waiting = db.query_one("SELECT COUNT(*) AS n FROM escalations WHERE status = 'answered' AND delivered_at IS NULL")
    return page(request, "dashboard.html", stages=stages, total=len(employees), opened=opened, states=states,
                places=places, recruiters=recruiters, red=red, topics=topics, model_up=llm.is_up(),
                waiting=waiting["n"])


@app.get("/recruiters/{recruiter_id}/qr.png")
def recruiter_qr(recruiter_id: int):
    recruiter = db.query_one("SELECT invite_code FROM recruiters WHERE id = ?", (recruiter_id,))
    return qr_response(engine.invite_link("hr_" + recruiter["invite_code"]) if recruiter else None)


@app.get("/knowledge")
def knowledge_page(request: Request):
    return page(request, "knowledge.html", documents=list_documents(), hits=None, question="", tags="")


def list_documents():
    return db.query("SELECT d.*, (SELECT COUNT(*) FROM chunks c WHERE c.document_id = d.id) AS chunks"
                    " FROM documents d ORDER BY d.title, d.id")


@app.post("/knowledge/upload")
def upload_document(request: Request, file: UploadFile | None = File(None), title: str = Form(""),
                    tags: str = Form("")):
    # A plain def: FastAPI runs it in a worker thread, so reading a big scan (OCR can take minutes)
    # does not stall the bot's event loop.
    if file is None or not file.filename:
        return back(request, "/knowledge", "Выберите файл.", bad=True)
    suffix = Path(file.filename).suffix.lower()
    if suffix not in UPLOAD_TYPES:
        return back(request, "/knowledge", "Такие файлы не принимаю: нужен PDF, DOCX, PPTX, TXT или MD.", bad=True)
    data = file.file.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        return back(request, "/knowledge", "Файл больше 20 МБ.", bad=True)

    folder = Path(cfg.DATA_DIR) / "uploads"
    # a random name: whatever the browser sent never reaches the file system
    path = folder / (uuid.uuid4().hex + suffix)
    try:
        folder.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    except OSError as e:
        log.warning("upload was not saved: %s", type(e).__name__)
        return back(request, "/knowledge", "Не удалось сохранить файл на сервере. Проверьте место и права на data/.",
                    bad=True)
    title = " ".join(title.split())[:MAX_NAME] or Path(file.filename).stem[:MAX_NAME]
    try:
        asyncio.run(knowledge.add_document(path, title, db.split_tags(tags)))
    except ValueError as e:
        path.unlink(missing_ok=True)
        return back(request, "/knowledge", str(e), bad=True)
    except Exception as e:
        # pdf, docx and pptx readers raise their own errors on a damaged file
        log.warning("document %s could not be read: %s", suffix, type(e).__name__)
        path.unlink(missing_ok=True)
        return back(request, "/knowledge", f"Не получилось прочитать файл: он повреждён или это не {suffix}.",
                    bad=True)
    return back(request, "/knowledge", "Документ добавлен.")


@app.post("/knowledge/{document_id}/delete")
def delete_document(request: Request, document_id: int):
    row = db.query_one("SELECT filename FROM documents WHERE id = ?", (document_id,))
    if row is None:
        return back(request, "/knowledge", "Этого документа уже нет.", bad=True)
    knowledge.delete_document(document_id)
    # only cabinet uploads live in data/uploads, documents loaded by tools/load_kb are somewhere else
    stored = Path(cfg.DATA_DIR) / "uploads" / Path(row["filename"]).name
    if stored.is_file():
        stored.unlink()
    return back(request, "/knowledge", "Документ удалён.")


@app.post("/knowledge/check")
async def check_question(request: Request, question: str = Form(""), tags: str = Form("")):
    question = question.strip()
    if not question:
        return back(request, "/knowledge", "Впишите вопрос.", bad=True)
    # a POST, so the question does not end up in the access log
    hits = await knowledge.search(question, db.split_tags(tags))
    return page(request, "knowledge.html", documents=list_documents(), hits=hits, question=question, tags=tags)


@app.get("/questions")
def questions_page(request: Request):
    rows = db.query(
        "SELECT q.*, e.full_name, e.tags, e.timezone AS employee_tz, e.workplace_id, w.name AS workplace_name,"
        " r.name AS recruiter_name, r.max_user_id AS recruiter_linked"
        " FROM escalations q JOIN employees e ON e.id = q.employee_id"
        " LEFT JOIN recruiters r ON r.id = e.recruiter_id LEFT JOIN workplaces w ON w.id = e.workplace_id"
        " ORDER BY q.id DESC"
    )
    now = datetime.now(ZoneInfo("UTC"))
    blocks: dict[str, list] = {}
    # a block per workplace, so each plant's questions are seen together
    for row in reversed([r for r in rows if r["status"] == "open"]):
        blocks.setdefault(row["workplace_name"] or "Место работы не указано", []).append(
            (row, read_context(row["context"]), hr.question_state(row, now)))
    waiting = []
    for row in rows:
        if row["status"] == "answered" and not row["delivered_at"]:
            employee = db.query_one("SELECT * FROM employees WHERE id = ?", (row["employee_id"],))
            later = hr.answer_time(employee)
            waiting.append((row, hours.when_text(later, employee["timezone"]) if later else "в ближайшую минуту"))
    answered = [row for row in rows if row["status"] != "open"][:30]
    took = {q["id"]: hours.waited_text(q["created_at"], datetime.fromisoformat(q["answered_at"]))
            for q in answered if q["answered_at"]}
    return page(request, "questions.html", blocks=blocks, answered=answered, waiting=waiting, took=took)


def read_context(raw: str) -> list[dict]:
    try:
        items = json.loads(raw)
    except ValueError:
        return []
    return [item for item in items if isinstance(item, dict)] if isinstance(items, list) else []


@app.post("/questions/{escalation_id}/answer")
async def answer_question(request: Request, escalation_id: int, answer: str = Form(""), employee_id: int = Form(0),
                          asked_at: str = Form("")):
    answer = answer.strip()
    # the form names the newcomer and the time too: after /delete the question number may go to a new question
    owner = db.query_one("SELECT e.max_user_id FROM escalations q JOIN employees e ON e.id = q.employee_id"
                         " WHERE q.id = ? AND q.employee_id = ? AND q.created_at = ?",
                         (escalation_id, employee_id, asked_at))
    if not answer:
        return back(request, "/questions", "Впишите ответ.", bad=True)
    if owner is not None and not owner["max_user_id"]:
        return back(request, "/questions", "Новичок не привязан к боту: удалил свои данные или не открыл "
                                           "ссылку. Ответить некому.", bad=True)
    try:
        later = await hr.deliver_answer(escalation_id, answer, "кабинет", employee_id, asked_at)
    except ValueError as e:
        return back(request, "/questions", str(e), bad=True)
    except Exception as e:
        # deliver_answer has already reopened the question, nothing is lost
        log.warning("escalation %s: answer was not sent: %s", escalation_id, type(e).__name__)
        return back(request, "/questions", "Не получилось отправить ответ в МАКС. Вопрос остался открытым, "
                                           "попробуйте ещё раз.", bad=True)
    if later is not None:
        employee = db.query_one("SELECT timezone FROM employees WHERE id = ?", (employee_id,))
        return back(request, "/questions", "Ответ принят. У новичка нерабочее время, ответ уйдёт "
                                           f"{hours.when_text(later, employee['timezone'])} по его времени.")
    return back(request, "/questions", "Ответ отправлен новичку.")


@app.post("/questions/{escalation_id}/send-now")
async def send_answer_now(request: Request, escalation_id: int):
    try:
        await hr.send_waiting_now(escalation_id)
    except ValueError as e:
        return back(request, "/questions", str(e), bad=True)
    except Exception as e:
        log.warning("escalation %s: waiting answer was not sent: %s", escalation_id, type(e).__name__)
        return back(request, "/questions", "Не получилось отправить ответ в МАКС, попробуйте ещё раз.", bad=True)
    return back(request, "/questions", "Ответ отправлен новичку.")


def read_scenario() -> tuple[list[Step], list[str]]:
    """The working scenario file: its steps and the problems that keep it from loading."""
    try:
        return load_scenario(cfg.SCENARIO_PATH), []
    except ScenarioError as e:
        return [], e.problems


@app.get("/scenario")
def scenario_page(request: Request):
    steps, problems = read_scenario()
    return page(request, "scenario.html", steps=len(steps), problems=problems, upload_problems=None)


@app.get("/scenario/download")
def download_scenario():
    if not Path(cfg.SCENARIO_PATH).is_file():
        raise HTTPException(404)
    return FileResponse(cfg.SCENARIO_PATH, filename="scenario.xlsx")


@app.post("/scenario/upload")
async def upload_scenario(request: Request, file: UploadFile | None = File(None)):
    if file is None or not file.filename:
        return back(request, "/scenario", "Выберите файл.", bad=True)
    if Path(file.filename).suffix.lower() != ".xlsx":
        return back(request, "/scenario", "Нужен файл Excel с расширением .xlsx.", bad=True)
    data = await file.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        return back(request, "/scenario", "Файл больше 20 МБ.", bad=True)

    working = Path(cfg.SCENARIO_PATH)
    # next to the working file: on the same disk the swap below is one rename, nobody sees half a file
    incoming = working.with_name(f".new-{uuid.uuid4().hex}.xlsx")
    try:
        incoming.write_bytes(data)
        problems = validate_scenario(incoming)
        if problems:
            steps, current = read_scenario()
            return page(request, "scenario.html", 400, steps=len(steps), problems=current, upload_problems=problems)
        if working.is_file():
            Path(cfg.DATA_DIR).mkdir(parents=True, exist_ok=True)
            shutil.copy2(working, Path(cfg.DATA_DIR) / "scenario-previous.xlsx")
        os.replace(incoming, working)
    except OSError as e:
        log.warning("scenario file was not written: %s", type(e).__name__)
        return back(request, "/scenario", "Не удалось записать файл сценария. Проверьте права на папку scenario/.",
                    bad=True)
    finally:
        incoming.unlink(missing_ok=True)

    problems = engine.reload_scenario()
    if problems:
        return back(request, "/scenario", "Файл заменён, но бот его не загрузил: " + problems[0], bad=True)
    return back(request, "/scenario", "Сценарий проверен, заменён и загружен в бота.")


@app.get("/export.xlsx")
def export_excel():
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, (header, rows) in (("Новички", newcomers_sheet()), ("Опросы", surveys_sheet()),
                                  ("Флаги", flags_sheet()), ("Вопросы", questions_sheet())):
        add_sheet(workbook, title, header, rows)
    out = io.BytesIO()
    workbook.save(out)
    return Response(out.getvalue(), media_type=XLSX_TYPE,
                    headers={"Content-Disposition": 'attachment; filename="korsik-report.xlsx"'})


def newcomers_sheet():
    rows = db.query(
        "SELECT e.*, r.name AS recruiter_name,"
        " (SELECT COUNT(*) FROM sent_steps s WHERE s.employee_id = e.id AND s.status != 'skipped') AS sent,"
        " (SELECT COUNT(*) FROM sent_steps s WHERE s.employee_id = e.id AND s.status = 'skipped') AS skipped,"
        " (SELECT COUNT(*) FROM answers a WHERE a.employee_id = e.id) AS answers,"
        " (SELECT COUNT(*) FROM escalations q WHERE q.employee_id = e.id AND q.status = 'open') AS open_questions,"
        " (SELECT COUNT(*) FROM pulse p WHERE p.employee_id = e.id AND p.risk = 2) AS red_flags,"
        " (SELECT name FROM workplaces w WHERE w.id = e.workplace_id) AS workplace_name"
        " FROM employees e LEFT JOIN recruiters r ON r.id = e.recruiter_id ORDER BY e.start_date, e.id"
    )
    header = ["ФИО", "Имя", "Место работы", "Дата выхода", "Часовой пояс", "Теги", "Рекрутер", "В боте с", "Обращение",
              "Шагов отправлено", "Шагов пропущено", "Ответов на опросы", "Открытых вопросов", "Красных флагов"]
    return header, [
        [e["full_name"], e["first_name"], e["workplace_name"] or "", e["start_date"], e["timezone"], e["tags"], e["recruiter_name"] or "",
         local_time(e["linked_at"]), {"ty": "на ты", "vy": "на вы"}.get(e["address"], ""), e["sent"], e["skipped"],
         e["answers"], e["open_questions"], e["red_flags"]]
        for e in rows
    ]


def surveys_sheet():
    rows = db.query("SELECT a.*, e.full_name FROM answers a JOIN employees e ON e.id = a.employee_id ORDER BY a.id")
    header = ["Новичок", "Шаг", "Вопрос", "Ответ", "Оценка", "Когда"]
    return header, [[a["full_name"], a["step_id"], a["question"], a["answer"], a["score"],
                     local_time(a["created_at"])] for a in rows]


def flags_sheet():
    rows = db.query("SELECT p.*, e.full_name FROM pulse p JOIN employees e ON e.id = p.employee_id"
                    " WHERE p.risk >= 1 ORDER BY p.id")
    header = ["Новичок", "Когда", "Риск", "Тема", "Тон", "Причина", "Сообщение"]
    return header, [[p["full_name"], local_time(p["created_at"]), RISK_NAMES[p["risk"]], p["topic"], p["tone"],
                     p["reason"], p["text"]] for p in rows]


def questions_sheet():
    rows = db.query("SELECT q.*, e.full_name, r.name AS recruiter_name FROM escalations q"
                    " JOIN employees e ON e.id = q.employee_id LEFT JOIN recruiters r ON r.id = e.recruiter_id"
                    " ORDER BY q.id")
    header = ["Новичок", "Рекрутер", "Когда", "Вопрос", "Статус", "Ответ", "Кто ответил", "Когда ответили",
              "Ответили за", "Напоминаний", "Доставлен новичку"]
    return header, [[q["full_name"], q["recruiter_name"] or "", local_time(q["created_at"]), q["question"],
                     "открыт" if q["status"] == "open" else "отвечен", q["answer"] or "", q["answered_by"] or "",
                     local_time(q["answered_at"]),
                     hours.waited_text(q["created_at"], datetime.fromisoformat(q["answered_at"]))
                     if q["answered_at"] else "", q["reminders"], local_time(q["delivered_at"])] for q in rows]


def add_sheet(workbook, title: str, header: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(header)
    for row in rows:
        sheet.append([clean_cell(value) for value in row])
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            # newcomers type this text: a value that starts with "=" would run as a formula in Excel
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    for column in sheet.columns:
        widest = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(widest + 2, 10), 60)


def clean_cell(value):
    if not isinstance(value, str):
        return value
    # control characters make openpyxl refuse the whole sheet, and a cell holds at most 32767 characters
    return ILLEGAL_CHARACTERS_RE.sub("", value)[:32000]
