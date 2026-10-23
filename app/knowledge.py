import logging
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path

import numpy as np
import pymupdf
import pytesseract
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from app import db, llm

log = logging.getLogger(__name__)

CHUNK_SIZE = 1000
TOP_K = 8
RRF_K = 60
EMBED_BATCH = 16
EMBED_PAUSE = timedelta(minutes=5)  # after a failed batch, before the next try
OCR_DPI = 300
HR_ANSWERS_TITLE = "Ответы HR"

# "6.3.3. Текст", "6.3.3 Текст" and, after OCR, "6.3.3.Текст"
CLAUSE_START = r"\d+(?:\.\d+)+(?:\.?\s|\.(?=[А-ЯЁA-Z]))"
CLAUSE = re.compile(r"(\d+(?:\.\d+)+)(?:\.?\s|\.(?=[А-ЯЁA-Z]))")
# OCR reads list dashes as hyphens, en or em dashes, underscores, sometimes with "=" after them
LIST_MARK = re.compile(r"^[-\u2013\u2014_][\s_=\u2013\u2014-]*(?=\w)")

DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
DRAWING = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

STOP_WORDS = {
    "что", "как", "где", "когда", "какой", "какая", "какое", "какие", "кто", "чем", "кому",
    "сколько", "почему", "зачем", "куда", "откуда", "для", "или", "это", "мне", "меня", "нас",
    "нам", "вас", "вам", "есть", "можно", "нужно", "надо", "при", "про", "под", "над", "без",
    "через", "его", "она", "они", "все", "всё", "так", "уже", "еще", "ещё", "только", "если",
    "тоже", "также", "этот", "эта", "эти", "тот", "там", "тут", "здесь", "нет", "привет",
    "пожалуйста", "подскажи", "подскажите", "скажи", "скажите", "расскажи", "расскажите",
}


@dataclass
class Hit:
    chunk_id: int
    title: str
    section: str
    text: str
    score: float
    keyword_rank: int | None
    vector_score: float | None


def extract_sections(path) -> list[tuple[str, str]]:
    """(section label, text) pairs: one per slide or page, a single one for plain text."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return [("", _read_text(path))]
    if suffix == ".pdf":
        return _pdf_sections(path)
    if suffix == ".docx":
        return [("", _docx_text(path))]
    if suffix == ".pptx":
        return _pptx_sections(path)
    raise ValueError(f"Файлы {suffix or 'без расширения'} не поддерживаются: нужен PDF, DOCX, PPTX, TXT или MD")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        # old HR files saved from Windows Notepad
        return raw.decode("cp1251", errors="replace")


def _pdf_sections(path: Path) -> list[tuple[str, str]]:
    sections = []
    with pymupdf.open(path) as pdf:
        for number, page in enumerate(pdf, 1):
            text = page.get_text().strip()
            # a scanned page has no text layer, at most a stray page number
            if len(text) < 20:
                text = tidy_ocr(ocr_page(page))
            if text:
                sections.append((f"стр. {number}", text))
    return sections


def ocr_page(page) -> str:
    pixmap = page.get_pixmap(dpi=OCR_DPI)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    try:
        return pytesseract.image_to_string(image, lang="rus")
    except (pytesseract.TesseractNotFoundError, pytesseract.TesseractError) as e:
        log.warning("ocr failed: %s", type(e).__name__)
        raise ValueError(
            "Это скан без текстового слоя, а распознавание (tesseract с русским языком) "
            "на сервере не работает. Загрузите текстовый вариант документа."
        ) from e


def tidy_ocr(text: str) -> str:
    """Glues wrapped OCR lines back into paragraphs and drops lines of pure noise."""
    paragraphs = []
    lines = []
    for line in text.splitlines() + [""]:
        line = line.strip()
        if not line:
            if lines:
                paragraphs.append(_join_lines(lines))
                lines = []
            continue
        if sum(ch.isalnum() for ch in line) < 2:
            continue
        line = LIST_MARK.sub("- ", line)
        if lines and (line.startswith("- ") or CLAUSE.match(line)):
            paragraphs.append(_join_lines(lines))
            lines = []
        lines.append(line)
    return "\n\n".join(paragraphs)


def _join_lines(lines: list[str]) -> str:
    text = lines[0]
    for line in lines[1:]:
        # a scan of a Word file rarely has soft hyphens, a line ending in "-" is usually "детей-инвалидов"
        if re.search(r"\w-$", text):
            text += line
        else:
            text += " " + line
    return text


def _docx_text(path: Path) -> str:
    parts = []
    for block in Document(path).iter_inner_content():
        if hasattr(block, "rows"):
            for row in block.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        elif block.text.strip():
            parts.append(block.text.strip())
    return "\n\n".join(parts)


def _pptx_sections(path: Path) -> list[tuple[str, str]]:
    sections = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        parts = _shape_texts(slide.shapes)
        # SmartArt text lives in a separate diagram part that python-pptx does not read
        for rel in slide.part.rels.values():
            if rel.reltype == RT.DIAGRAM_DATA:
                parts.append(_smartart_text(rel.target_part.blob))
        text = "\n\n".join(part for part in parts if part)
        if text:
            sections.append((f"слайд {number}", text))
    return sections


def _shape_texts(shapes) -> list[str]:
    texts = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts += _shape_texts(shape.shapes)
        elif shape.has_text_frame:
            texts.append(shape.text_frame.text.replace("\v", "\n").strip())
        elif shape.has_table:
            for row in shape.table.rows:
                texts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return texts


def _smartart_text(xml: bytes) -> str:
    root = ET.fromstring(xml)
    texts = {}
    top = None
    for point in root.iter(DGM + "pt"):
        kind = point.get("type", "node")
        if kind == "doc":
            top = point.get("modelId")
        elif kind in ("node", "asst"):
            paragraphs = ("".join(t.text or "" for t in p.iter(DRAWING + "t")) for p in point.iter(DRAWING + "p"))
            texts[point.get("modelId")] = " ".join(p.strip() for p in paragraphs if p.strip())

    children = {}
    for link in root.iter(DGM + "cxn"):
        if link.get("type", "parOf") == "parOf":
            children.setdefault(link.get("srcId"), []).append((int(link.get("srcOrd", 0)), link.get("destId")))

    # the point list is not in reading order, the parent links are
    blocks = []

    def walk(node, depth):
        for _, child in sorted(children.get(node, [])):
            text = texts.get(child)
            if text and depth > 0 and blocks:
                blocks[-1] += "\n- " + text
            elif text:
                blocks.append(text)
            walk(child, depth + 1)

    walk(top, 0)
    if not blocks:
        blocks = [text for text in texts.values() if text]
    return "\n\n".join(blocks)


def split_chunks(sections) -> list[tuple[str, str]]:
    """Numbered clauses keep their number ("п. 6.3.3"), other text goes by paragraphs up to CHUNK_SIZE."""
    chunks = []
    for label, text in sections:
        groups = []
        headings = []
        for paragraph in _paragraphs(text):
            if _is_heading(paragraph):
                headings.append(paragraph)
                continue
            number = _clause_number(paragraph) or next(
                (n for n in map(_clause_number, reversed(headings)) if n), None
            )
            if number or not groups:
                groups.append([f"п. {number}" if number else label, headings + [paragraph]])
            else:
                groups[-1][1] += headings + [paragraph]
            headings = []
        if headings:
            if groups:
                groups[-1][1] += headings
            else:
                groups.append([label, headings])
        for section, paragraphs in groups:
            chunks += [(section, piece) for piece in _pack(paragraphs)]
    return chunks


def _paragraphs(text: str) -> list[str]:
    parts = re.split(r"\n\s*\n|\n(?=" + CLAUSE_START + ")", text.replace("\r\n", "\n"))
    return [part.strip() for part in parts if part.strip()]


def _clause_number(paragraph: str) -> str | None:
    match = CLAUSE.match(paragraph)
    return match.group(1) if match else None


def _is_heading(paragraph: str) -> bool:
    return len(paragraph) <= 80 and "\n" not in paragraph and not paragraph.endswith((".", ",", ";", ":", "!", "?"))


def _pack(paragraphs: list[str], size: int = CHUNK_SIZE) -> list[str]:
    pieces = []
    current = ""
    for paragraph in paragraphs:
        for part in _cut(paragraph, size):
            if current and len(current) + 2 + len(part) > size:
                pieces.append(current)
                current = part
            else:
                current = f"{current}\n\n{part}" if current else part
    if current:
        pieces.append(current)
    return pieces


def _cut(text: str, size: int) -> list[str]:
    parts = []
    while len(text) > size:
        cut = max(text.rfind(". ", 0, size), text.rfind("; ", 0, size), text.rfind("\n", 0, size))
        if cut < size // 2:
            cut = text.rfind(" ", 0, size)
        if cut <= 0:
            cut = size - 1
        parts.append(text[: cut + 1].strip())
        text = text[cut + 1 :].strip()
    parts.append(text)
    return parts


def _fts_text(text: str) -> str:
    # unicode61 folds case for Cyrillic but treats ё and е as different letters
    return text.replace("ё", "е").replace("Ё", "Е")


def _insert_chunks(conn, document_id: int, chunks: list[tuple[str, str]]) -> None:
    for section, text in chunks:
        chunk_id = conn.execute(
            "INSERT INTO chunks (document_id, section, text) VALUES (?, ?, ?)", (document_id, section, text)
        ).lastrowid
        conn.execute("INSERT INTO chunks_fts (text, chunk_id) VALUES (?, ?)", (_fts_text(text), chunk_id))


def _insert_document(title: str, filename: str, tags: list[str], kind: str, chunks) -> int:
    conn = db.connect()
    try:
        with conn:
            document_id = conn.execute(
                "INSERT INTO documents (title, filename, tags, kind, created_at) VALUES (?, ?, ?, ?, ?)",
                (title, filename, ", ".join(tags), kind, db.now_iso()),
            ).lastrowid
            _insert_chunks(conn, document_id, chunks)
        return document_id
    finally:
        conn.close()


async def add_document(path, title: str, tags: list[str], kind: str = "file") -> int:
    """Stores a file as a document. A file with the same name replaces the old version."""
    path = Path(path)
    chunks = split_chunks(extract_sections(path))
    if not chunks:
        raise ValueError("В документе не нашлось текста")
    for row in db.query("SELECT id FROM documents WHERE filename = ?", (path.name,)):
        delete_document(row["id"])
    document_id = _insert_document(title, path.name, tags, kind, chunks)
    log.info("document %s added: %d chunks", document_id, len(chunks))
    if llm.is_up():
        await embed_missing()
    return document_id


def delete_document(doc_id: int) -> None:
    conn = db.connect()
    try:
        with conn:
            conn.execute(
                "DELETE FROM chunks_fts WHERE chunk_id IN (SELECT id FROM chunks WHERE document_id = ?)", (doc_id,)
            )
            conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
    finally:
        conn.close()


async def embed_missing() -> int:
    """Computes vectors for chunks stored while the model was away. Returns how many were done."""
    # the watch loop calls this every minute, and a batch that timed out would only time out again
    if (db.kv_get("embed_retry_at") or "") > db.now_iso():
        return 0
    done = 0
    while True:
        rows = db.query(
            "SELECT c.id, c.text, d.title FROM chunks c JOIN documents d ON d.id = c.document_id"
            " WHERE c.embedding IS NULL ORDER BY c.id LIMIT ?",
            (EMBED_BATCH,),
        )
        if not rows:
            if done:
                log.info("%d chunks embedded", done)
            return done
        try:
            vectors = await llm.embed([f"{row['title']}\n{row['text']}" for row in rows], mark_down=False)
        except llm.LLMUnavailable:
            retry_at = datetime.now(timezone.utc) + EMBED_PAUSE
            db.kv_set("embed_retry_at", retry_at.isoformat(timespec="seconds"))
            log.info("embeddings postponed, %d done so far", done)
            return done
        conn = db.connect()
        try:
            with conn:
                for row, vector in zip(rows, vectors):
                    blob = np.asarray(vector, dtype=np.float32).tobytes()
                    conn.execute("UPDATE chunks SET embedding = ? WHERE id = ?", (blob, row["id"]))
        finally:
            conn.close()
        done += len(rows)


def add_hr_answer(question: str, answer: str, tags: list[str]) -> None:
    """Adds a pair to the "Ответы HR" document with the same tags. The vector comes later from embed_missing()."""
    wanted = {tag.lower() for tag in tags}
    document = next(
        (
            row
            for row in db.query("SELECT id, tags FROM documents WHERE kind = 'hr_answers' ORDER BY id")
            if {tag.lower() for tag in db.split_tags(row["tags"])} == wanted
        ),
        None,
    )
    chunk = ("", f"Вопрос: {question.strip()}\nОтвет: {answer.strip()}")
    if document is None:
        _insert_document(HR_ANSWERS_TITLE, "", tags, "hr_answers", [chunk])
        return
    conn = db.connect()
    try:
        with conn:
            _insert_chunks(conn, document["id"], [chunk])
    finally:
        conn.close()


def documents_for(employee_tags) -> list:
    """Documents without tags plus documents sharing at least one tag with the employee."""
    if isinstance(employee_tags, str):
        employee_tags = db.split_tags(employee_tags)
    wanted = {tag.lower() for tag in employee_tags}
    visible = []
    for row in db.query("SELECT * FROM documents ORDER BY title, id"):
        tags = {tag.lower() for tag in db.split_tags(row["tags"])}
        if not tags or tags & wanted:
            visible.append(row)
    return visible


def _match_query(query: str) -> str:
    terms = []
    for word in re.findall(r"[^\W_]+", _fts_text(query.lower())):
        if len(word) < 3 or word in STOP_WORDS:
            continue
        terms.append(f'"{word[: max(4, len(word) - 2)]}"*')
    return " OR ".join(dict.fromkeys(terms))


def _keyword_search(query: str, doc_ids: list[int]) -> list[int]:
    match = _match_query(query)
    if not match:
        return []
    marks = ",".join("?" * len(doc_ids))
    rows = db.query(
        "SELECT f.chunk_id FROM chunks_fts f JOIN chunks c ON c.id = f.chunk_id"
        f" WHERE chunks_fts MATCH ? AND c.document_id IN ({marks})"
        " ORDER BY bm25(chunks_fts) LIMIT ?",
        (match, *doc_ids, TOP_K),
    )
    return [row["chunk_id"] for row in rows]


async def _vector_search(query: str, doc_ids: list[int]) -> list[tuple[int, float]]:
    target = np.asarray((await llm.embed([query]))[0], dtype=np.float32)
    marks = ",".join("?" * len(doc_ids))
    rows = db.query(
        f"SELECT id, embedding FROM chunks WHERE embedding IS NOT NULL AND document_id IN ({marks})", doc_ids
    )
    ids, vectors = [], []
    for row in rows:
        vector = np.frombuffer(row["embedding"], dtype=np.float32)
        # vectors from another embedding model have another size, they wait for a reload
        if vector.shape == target.shape:
            ids.append(row["id"])
            vectors.append(vector)
    if not ids:
        return []
    matrix = np.vstack(vectors)
    norms = np.linalg.norm(matrix, axis=1) * np.linalg.norm(target)
    scores = matrix @ target / np.where(norms == 0, 1, norms)
    best = np.argsort(-scores)[:TOP_K]
    return [(ids[i], float(scores[i])) for i in best]


async def search(query: str, employee_tags: list[str], limit: int = 5, use_vectors: bool = True) -> list[Hit]:
    doc_ids = [row["id"] for row in documents_for(employee_tags)]
    if not doc_ids:
        return []

    keyword = _keyword_search(query, doc_ids)
    vector = []
    if use_vectors and llm.is_up():
        try:
            vector = await _vector_search(query, doc_ids)
        except llm.LLMUnavailable:
            log.info("vector search skipped, model is unavailable")

    # reciprocal rank fusion: only positions matter, so bm25 and cosine need no common scale
    scores = {}
    for rank, chunk_id in enumerate(keyword, 1):
        scores[chunk_id] = scores.get(chunk_id, 0) + 1 / (RRF_K + rank)
    for rank, (chunk_id, _) in enumerate(vector, 1):
        scores[chunk_id] = scores.get(chunk_id, 0) + 1 / (RRF_K + rank)
    best = sorted(scores, key=scores.get, reverse=True)[:limit]
    if not best:
        return []

    keyword_rank = {chunk_id: rank for rank, chunk_id in enumerate(keyword, 1)}
    vector_score = dict(vector)
    marks = ",".join("?" * len(best))
    rows = {
        row["id"]: row
        for row in db.query(
            "SELECT c.id, c.section, c.text, d.title FROM chunks c JOIN documents d ON d.id = c.document_id"
            f" WHERE c.id IN ({marks})",
            best,
        )
    }
    return [
        Hit(
            chunk_id=chunk_id,
            title=rows[chunk_id]["title"],
            section=rows[chunk_id]["section"],
            text=rows[chunk_id]["text"],
            score=scores[chunk_id],
            keyword_rank=keyword_rank.get(chunk_id),
            vector_score=vector_score.get(chunk_id),
        )
        for chunk_id in best
    ]
