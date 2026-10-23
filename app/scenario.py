import calendar
import logging
import re
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

log = logging.getLogger(__name__)

SHEET = "Сценарий"
COLUMNS = {
    "шаг": "id",
    "когда": "when",
    "после": "after",
    "теги": "tags",
    "тип": "kind",
    "текст на «ты»": "text_ty",
    "текст на «вы»": "text_vy",
    "кнопки": "buttons",
    "плохой ответ": "bad_answer",
    "картинка": "image",
}
REQUIRED_COLUMNS = ("id", "when", "kind", "text_ty")
KINDS = ("сообщение", "обращение", "вопрос", "опрос", "шкала", "текст")
MAX_TEXT = 4000
IMAGE_RE = re.compile(r"[\w.-]+\.(png|jpe?g|gif)", re.IGNORECASE)

START = "старт"
# Steps that never go out on schedule: the engine takes their text when it needs it.
SPECIAL_WHEN = ("напоминание", "эскалация")

WHEN_RE = re.compile(
    r"(?:(?P<days>[+-]\d+)д|р(?P<workday>[1-9]\d*)|м(?P<months>[1-9]\d*)(?:(?P<extra>[+-]\d+)д)?|(?P<end>конец))"
    r" (?P<hour>\d{1,2})[:.](?P<minute>\d{2})"
)
HR_NOTE_RE = re.compile(r"\[HR:[^\]]*\]")


@dataclass
class Button:
    text: str
    url: str | None = None


@dataclass
class Step:
    id: str
    when: str
    after: tuple[str, str] | None
    tags: list[str]
    kind: str
    text_ty: str
    text_vy: str
    buttons: list[Button]
    bad_answer: str
    image: str
    row: int


class ScenarioError(Exception):
    def __init__(self, problems: list[str]):
        super().__init__("\n".join(problems))
        self.problems = problems


def load_scenario(path) -> list[Step]:
    steps, problems = _read(path)
    if problems:
        raise ScenarioError(problems)
    return steps


def validate_scenario(path) -> list[str]:
    return _read(path)[1]


def load_holidays(path) -> set[date]:
    path = Path(path)
    if not path.is_file():
        return set()
    days = set()
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.split("#", 1)[0].strip()
        if not line:
            continue
        try:
            days.add(date.fromisoformat(line))
        except ValueError:
            log.warning("holidays: skipped bad line %r", line)
    return days


def tags_match(step_tags: list[str], employee_tags: list[str]) -> bool:
    have = {t.strip().lower() for t in employee_tags if t.strip()}
    wanted, banned = set(), set()
    for tag in step_tags:
        tag = tag.strip().lower()
        if tag.startswith("-"):
            banned.add(tag[1:].strip())
        elif tag:
            wanted.add(tag)
    if have & banned:
        return False
    return not wanted or bool(have & wanted)


def step_due_at(step: Step, start_date: date, tz: str, holidays: set[date],
                linked_at: datetime | None, probation_months: int = 3) -> datetime | None:
    if step.after:
        return None
    when = _norm(step.when)
    if when == START:
        if linked_at is None:
            return None
        if linked_at.tzinfo is None:  # the db keeps UTC without an offset
            linked_at = linked_at.replace(tzinfo=timezone.utc)
        return linked_at.astimezone(timezone.utc)
    parsed = _parse_when(when)
    if parsed is None:
        return None
    match, at = parsed
    day = _due_date(match, start_date, holidays, probation_months)
    return datetime.combine(day, at, tzinfo=ZoneInfo(tz)).astimezone(timezone.utc)


def render_text(step: Step, address: str | None, first_name: str) -> str:
    # until the employee picks "ты" or "вы" the bot talks on "вы" (spec 4.4)
    text = step.text_ty if address == "ty" else (step.text_vy or step.text_ty)
    text = HR_NOTE_RE.sub("", text)
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text.replace("{имя}", first_name)


def is_bad_answer(step: Step, answer: str, score: int | None) -> bool:
    rule = step.bad_answer.strip()
    if not rule:
        return False
    if rule.startswith("<"):
        if score is None:
            try:
                score = int(answer.strip())
            except ValueError:
                return False
        return score < int(rule[1:])
    bad = {option.strip().lower() for option in rule.split("|")}
    return answer.strip().lower() in bad


def scale_buttons() -> list[list[Button]]:
    return [[Button(str(n)) for n in range(1, 6)], [Button(str(n)) for n in range(6, 11)]]


def _read(path) -> tuple[list[Step], list[str]]:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # zip, xml and format errors all mean the same thing for HR
        return [], [f"не удалось открыть файл: {e}"]
    if SHEET not in wb.sheetnames:
        return [], [f"нет листа «{SHEET}»"]
    rows = list(wb[SHEET].iter_rows(values_only=True))
    if not rows:
        return [], [f"лист «{SHEET}» пустой"]

    positions = {}
    for i, title in enumerate(rows[0]):
        key = COLUMNS.get(_cell(title).lower())
        if key:
            positions[key] = i
    missing = [title for title, key in COLUMNS.items() if key in REQUIRED_COLUMNS and key not in positions]
    if missing:
        return [], ["нет колонок: " + ", ".join(f"«{title}»" for title in missing)]

    steps, problems = [], []
    for row, values in enumerate(rows[1:], start=2):
        cells = dict.fromkeys(COLUMNS.values(), "")
        cells.update({key: _cell(values[i]) for key, i in positions.items() if i < len(values)})
        if not any(cells[key] for key in ("id", "when", "after", "kind", "text_ty")):
            continue
        steps.append(_parse_row(cells, row, problems))

    by_id = {}
    for step in steps:
        if step.id in by_id:
            problems.append(f"строка {step.row}, шаг {step.id}: повтор номера шага "
                            f"(первый раз в строке {by_id[step.id].row})")
        else:
            by_id[step.id] = step
    for step in steps:
        if not step.after:
            continue
        target_id, button = step.after
        target = by_id.get(target_id)
        if target is None:
            problems.append(f"строка {step.row}, шаг {step.id}: в «После» указан шаг {target_id}, такого нет")
        elif button.lower() not in {b.text.lower() for b in target.buttons if not b.url}:
            problems.append(f"строка {step.row}, шаг {step.id}: у шага {target_id} нет кнопки «{button}»")
    return steps, problems


def _parse_row(cells: dict[str, str], row: int, problems: list[str]) -> Step:
    where = f"строка {row}, шаг {cells['id']}" if cells["id"] else f"строка {row}"

    def problem(text):
        problems.append(f"{where}: {text}")

    if not cells["id"]:
        problem("нет номера шага")

    kind = cells["kind"].lower()
    if kind not in KINDS:
        problem(f"неизвестный тип «{cells['kind']}», можно: {', '.join(KINDS)}")

    when = _norm(cells["when"])
    after = None
    if cells["after"]:
        target, _, button = cells["after"].partition(":")
        if target.strip() and button.strip():
            after = (target.strip(), button.strip())
        else:
            problem(f"«После» надо писать как «25: Да!», а написано «{cells['after']}»")
    if when and cells["after"]:
        problem("заполнены и «Когда», и «После»: нужно что-то одно")
    elif not when and not cells["after"]:
        problem("не заполнены ни «Когда», ни «После»")
    elif when and when != START and when not in SPECIAL_WHEN and _parse_when(when) is None:
        problem(f"не понял «Когда»: «{cells['when']}»")

    tags = [t.strip() for t in cells["tags"].split(",") if t.strip()]
    if any(t.lstrip("-").strip() == "" for t in tags):
        problem("пустой тег после минуса")

    buttons = []
    for part in cells["buttons"].split("|"):
        part = part.strip()
        if not part:
            continue
        if "=" in part:
            text, _, url = (s.strip() for s in part.partition("="))
            if not url.startswith(("http://", "https://")):
                problem(f"у кнопки «{text}» ссылка должна начинаться с http:// или https://")
            buttons.append(Button(text, url))
        else:
            buttons.append(Button(part))
    answers = [b.text for b in buttons if not b.url]

    if kind == "обращение" and len(answers) != 2:
        problem("у обращения должно быть две кнопки: на «ты» и на «вы»")
    elif kind in ("вопрос", "опрос") and not answers:
        problem(f"у шага типа «{kind}» нет кнопок")
    elif kind == "шкала" and buttons:
        problem("у шкалы кнопки 1-10 ставятся сами, колонку «Кнопки» оставь пустой")
    elif kind in ("сообщение", "текст") and answers:
        problem(f"у шага типа «{kind}» могут быть только кнопки-ссылки")

    bad = cells["bad_answer"]
    if bad:
        if kind == "шкала":
            if not re.fullmatch(r"<\s*(10|[2-9])", bad):
                problem(f"плохой ответ для шкалы пишется как «<8», а написано «{bad}»")
            bad = bad.replace(" ", "")
        elif kind == "опрос":
            known = {a.lower() for a in answers}
            for option in bad.split("|"):
                if option.strip().lower() not in known:
                    problem(f"плохой ответ «{option.strip()}» не совпадает ни с одной кнопкой")
        else:
            problem(f"«Плохой ответ» бывает только у опроса и шкалы, а тут тип «{kind}»")

    image = cells["image"]
    if image and not IMAGE_RE.fullmatch(image):
        problem(f"картинка «{image}»: нужно имя файла из data/images, например img_03.png")

    step = Step(id=cells["id"], when="" if after else when, after=after, tags=tags, kind=kind,
                text_ty=cells["text_ty"], text_vy=cells["text_vy"], buttons=buttons,
                bad_answer=bad, image=image, row=row)

    for label, text in (("ты", step.text_ty), ("вы", step.text_vy)):
        if len(text) > MAX_TEXT:
            problem(f"текст на «{label}» длиннее {MAX_TEXT} символов ({len(text)})")
    if not render_text(step, "ty", "Имя") or not render_text(step, "vy", "Имя"):
        problem("пустой текст (строки [HR: ...] новичку не уходят)")
    return step


def _cell(value) -> str:
    # Excel turns "25" into 25 and "25.1" into 25.1; ids and texts are strings for us
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _norm(when: str) -> str:
    return " ".join(str(when or "").split()).lower()


def _parse_when(when: str) -> tuple[re.Match, time] | None:
    match = WHEN_RE.fullmatch(when)
    if not match:
        return None
    try:
        return match, time(int(match["hour"]), int(match["minute"]))
    except ValueError:
        return None


def _due_date(match: re.Match, start: date, holidays: set[date], probation_months: int) -> date:
    if match["workday"]:
        day = _next_working(start, holidays, 1)
        for _ in range(int(match["workday"]) - 1):
            day = _next_working(day + timedelta(days=1), holidays, 1)
        return day
    if match["days"]:
        day = start + timedelta(days=int(match["days"]))
    elif match["months"]:
        day = _add_months(start, int(match["months"])) + timedelta(days=int(match["extra"] or 0))
    else:
        day = _add_months(start, probation_months) - timedelta(days=1)
    # before the first day we move back so the employee is not late, after it we move forward
    return _next_working(day, holidays, -1 if day < start else 1)


def _next_working(day: date, holidays: set[date], direction: int) -> date:
    while day.weekday() >= 5 or day in holidays:
        day += timedelta(days=direction)
    return day


def _add_months(day: date, months: int) -> date:
    month_index = day.month - 1 + months
    year, month = day.year + month_index // 12, month_index % 12 + 1
    return date(year, month, min(day.day, calendar.monthrange(year, month)[1]))
